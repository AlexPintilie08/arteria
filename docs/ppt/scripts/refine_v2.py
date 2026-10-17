"""V2 refinements and semantic native groups for PowerPoint animations."""
from contextlib import contextmanager
from types import SimpleNamespace
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches


def refine(slides,copy,api):
    b=SimpleNamespace(**api);plan=[]
    @contextmanager
    def beat(s,name,trigger='click',effect='fade',duration=.4,delay=0):
        old={sh.shape_id for sh in s.shapes}
        yield
        selected=[sh for sh in s.shapes if sh.shape_id not in old]
        group(s,name,selected,trigger,effect,duration,delay)
    def group(s,name,shapes,trigger='click',effect='fade',duration=.4,delay=0):
        if not shapes:return
        g=s.shapes.add_group_shape(shapes);g.name='reveal:'+name
        plan.append({'slide':list(slides).index(s)+1,'shape':g.name,'trigger':trigger,
                     'effect':effect,'duration':duration,'delay':delay})
        return g
    def matching(s,*prefixes):
        return [sh for sh in s.shapes if sh.has_text_frame and any(sh.text.startswith(p) for p in prefixes)]
    def addtext(*a,**kw):return b.text(*a,**kw)

    # Problem: orthogonal branches, matching category proportions.
    s=slides[1];b.base(s,2,'P / PROBLEMA',copy[2][0],source='Situație propusă · README.md și docs/ARCHITECTURE.md, audit 2026')
    addtext(s,copy[2][1],.65,2.12,12,.55,25,b.INK,True)
    addtext(s,copy[2][2],.65,2.82,12,.55,23)
    with beat(s,'s02_intervention','auto'):
        b.node(s,'Intervenție','',4.88,3.65,3.57,.88)
    for j,(x,label) in enumerate(zip([.65,3.78,6.91,10.04],['Transport public','Pietoni','Biciclete','Trafic auto'])):
        with beat(s,f's02_category_{j+1}','click' if j%2==0 else 'after'):
            # Each branch runs to a common horizontal level, without diagonals.
            b.link(s,6.665,4.53,6.665,4.92,'718C9C')
            b.link(s,6.665,4.92,x+1.32,4.92,'718C9C')
            b.link(s,x+1.32,4.92,x+1.32,5.13,'718C9C')
            addtext(s,label,x,5.33,2.64,.6,22,b.INK,True,b.PP_ALIGN.CENTER)
    addtext(s,copy[2][4],.65,6.17,12,.55,25)
    addtext(s,copy[2][5],.65,6.72,12,.3,14,b.SECOND)

    # Alternatives: independent native table rows allow a meaningful sequence.
    s=slides[2];b.base(s,3,'R / ALTERNATIVE',copy[3][0],source=copy[3][-1])
    for j,row in enumerate([['QGIS','Date și analiză geografică'],['SUMO','Simulare microscopică de trafic'],['PMUD Iași','Context de planificare locală']]):
        with beat(s,f's03_row_{j+1}'):
            b.table(s,[row],.65,2.18+j*.94,[3,9.03],.85)
    with beat(s,'s03_value'):
        addtext(s,'ARTERIA propune un parcurs comun:',.65,5.35,12,.5,25,b.INK,True)
        addtext(s,'zonă → alternative → comparație',.65,5.99,12,.5,28,'315B70')
    addtext(s,copy[3][-2],.65,6.62,12,.35,16,b.SECOND)

    s=slides[3];b.base(s,4,'B / SOLUȚIA',copy[4][0],True)
    for j,(x,label) in enumerate(zip([.65,4.83,9.01],['Alege zona','Examinează\nalternativele','Compară\nscenariul'])):
        with beat(s,f's04_stage_{j+1}'):
            addtext(s,f'0{j+1}',x,2.65,3.6,.85,48,b.BLUE,True)
            addtext(s,label,x,3.72,3.6,1.25,27,b.PALE,True)
        if j<2:
            with beat(s,f's04_link_{j+1}','after','wipe',.3):b.link(s,x+3.15,3.13,x+3.8,3.13)
    addtext(s,copy[4][-2],.65,5.54,12,.6,24,b.AMBER)
    addtext(s,copy[4][-1],.65,6.24,12,.6,24,b.PALE)

    # Backend has a wider central block for real, explicitly demo endpoints.
    s=slides[5];b.base(s,6,'B / CONSTRUCȚIE',copy[6][0],source='backend/main.py · backend/server.py · sim/ · Audit 2026; legături întrerupte = integrări urbane viitoare')
    for name,title,body,x,w in [('frontend','Frontend','Next.js / MapLibre\nFixtures locale',.65,3.45),('backend','Backend','FastAPI: GET /health\nDemo HTTP: /analyze\n/propose · /evaluate',4.6,4.15),('sim','Sim','NetworkX sintetic\nSmoke SUMO opțional',9.25,3.43)]:
        with beat(s,'s06_'+name):b.node(s,title,body,x,2.27,w,2.5,size=20)
        if name!='frontend':
            with beat(s,'s06_link_'+name,'after','wipe',.3):
                x1,x2=(4.1,4.6) if name=='backend' else (8.75,9.25)
                b.link(s,x1,3.48,x2,3.48,'607B8B',True)
    addtext(s,'Endpoint-uri demo: date fixe, fără LLM sau evaluare SUMO.',.65,4.96,12,.45,18,b.SECOND,True)
    with beat(s,'s06_ai_future'):
        b.link(s,8.5,4.77,8.5,5.6,'607B8B',True)
        b.node(s,'AI planificat','Alternative structurate; validare înainte de evaluare',2.1,5.6,9.15,1.22,size=20)
    # Shorter horizontal node: separate body and title for comfortable fit.
    ai=s.shapes[-1]
    for sh in ai.shapes:
        if sh.has_text_frame and sh.text.startswith('Alternative structurate'):
            sh.top=Inches(6.21);sh.height=Inches(.42)

    # Balanced, equally styled team; no inferred ownership of particular files.
    s=slides[8];b.base(s,9,'O / ECHIPA','ARION',True,'Responsabilități confirmate de echipă · Cerința V2, 18 octombrie 2026')
    profiles=[('Iustin Fagateanu','Frontend and\nweb experience'),('Alex Pintilie','Bootstrap and\nsimulation foundation'),('Fagateanu Eusebiu','Backend')]
    for j,(name,role) in enumerate(profiles):
        with beat(s,f's09_member_{j+1}','click' if j==0 else 'after',duration=.4,delay=0 if j==0 else .06):
            x=.65+j*4.14
            b.node(s,name,role,x,2.65,3.75,2.55,True,size=22)
            # Names remain the same readable size for every profile.
            g=s.shapes[-1] if s.shapes[-1].shape_type==6 else None
            # Grouping occurs on context exit; text adjustments happen below.
    for g in s.shapes:
        if g.shape_type==6 and g.name.startswith('reveal:s09'):
            for sh in g.shapes:
                if sh.has_text_frame and sh.text in [p[0] for p in profiles]:
                    for p in sh.text_frame.paragraphs:
                        for run in p.runs:run.font.size=b.Pt(25)
                    sh.height=Inches(.65)
    addtext(s,'Următoarea nevoie: validare cu un specialist în mobilitate',.65,5.92,12,.8,24,b.PALE)

    # AI annex: consistent coordinates, one subtle automatic diagram reveal.
    s=slides[11];b.base(s,12,'ANEXA B / AI',copy[12][0],True,copy[12][-1])
    with beat(s,'s12_proposed_flow','auto',duration=.35):
        for j,(x,title,body) in enumerate([(.65,'Dovezi','și catalog'),(3.8,'LLM','Alternative'),(6.95,'Reguli','și geometrie'),(10.1,'Evaluare','independentă')]):
            b.node(s,title,body,x,2.65,2.58,1.8,True,size=20)
            if j<3:b.link(s,x+2.58,3.55,x+3.15,3.55,b.BLUE,True)
    addtext(s,copy[12][2],.65,4.86,12,.6,27,b.AMBER,True)
    addtext(s,copy[12][3],.65,5.68,12,.65,23,b.PALE)
    addtext(s,copy[12][4],.65,6.45,12,.55,25,b.BLUE)

    # Semantic groups for unchanged slides; disclosures intentionally stay static.
    s=slides[0]
    group(s,'s01_title',matching(s,'ARTERIA')[-1:],'auto',duration=.4)
    group(s,'s01_subtitle',matching(s,'Planificare urbană','Compararea intervențiilor'),'after',duration=.4)
    group(s,'s01_product',[sh for sh in s.shapes if sh.shape_type==13],'after',duration=.45)
    s=slides[4]
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text.startswith('Selectează reperul'):sh._element.getparent().remove(sh._element)
    group(s,'s05_product',[sh for sh in s.shapes if sh.shape_type==13],'auto',duration=.45)
    for j,label in enumerate(['Selectează reperul','Alege o alternativă','Deschide comparația']):
        with beat(s,f's05_step_{j+1}'):
            addtext(s,f'0{j+1}',.65,2.27+j*.88,.5,.45,17,b.BLUE,True)
            addtext(s,label,1.25,2.24+j*.88,3.45,.65,22,b.PALE,True)
    s=slides[6]
    group(s,'s07_customer',matching(s,'Cumpărător vizat','Administrație locală'))
    group(s,'s07_price',matching(s,'1.500 €','/ lună'))
    group(s,'s07_value',matching(s,'Valoare de verificat:','Costuri de măsurat:'))
    s=slides[7]
    group(s,'s08_first_case',matching(s,'Iași: primul'))
    group(s,'s08_outreach',matching(s,'Apoi:','Canal propus:'))
    group(s,'s08_commercial',matching(s,'Exercițiu comercial','10 ×'))
    s=slides[9]
    for j,x in enumerate([.65,4.83,9.01]):
        shapes=[sh for sh in s.shapes if sh.has_text_frame and abs(sh.left/914400-x)<.02 and 2.1<sh.top/914400<3.4]
        group(s,f's10_priority_{j+1}',shapes)
    group(s,'s10_request',matching(s,'Cererea noastră','O introducere către'))
    s=slides[10];group(s,'s11_test_result',matching(s,'5 / 5','teste ale clientului API'),'auto',duration=.35)
    s=slides[12];group(s,'s13_method',[sh for sh in s.shapes if sh.has_table],'auto',duration=.35)

    # Explicitly suppress theme-inherited shadows on native diagram objects.
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type==6:yield from walk(sh.shapes)
            else:yield sh
    for s in slides:
        for sh in walk(s.shapes):
            sp=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
            if sp is not None and sp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst') is None:
                sp.append(OxmlElement('a:effectLst'))
    return sorted(plan,key=lambda p:p['slide'])
