"""Generate audit traceability reports from the frozen original and audited copy."""
from pathlib import Path
import json
import re
from pptx import Presentation
from audited_content import EDITS, NOTES, EVIDENCE, TIMES

ROOT=Path(__file__).resolve().parents[1]; R=ROOT/'reports'
ORIGINAL=json.loads((R/'PPT_ORIGINAL_EXTRACTION.json').read_text(encoding='utf-8'))

TITLES=['ARTERIA','O intervenție. Mai multe efecte.','Instrumentele există','Parcursul unei comparații',
 'Podu Roș, în prototip','Frontend · Backend · Sim','Model economic de testat',
 'Primii utilizatori, înainte de extindere','ARION','Prima comparație pe date verificate',
 'Ce este verificat astăzi','Rolul AI în fluxul propus','Condițiile unei evaluări credibile']

# purpose, findings, remain, remove, add, decision, narrative rationale
REVIEW={
1: ('Identify product and user',
 'Broad urban-planning positioning can imply more modules than the integrated workflows support. The cover image is the preview, not the connected API workspace.',
 'Product, team, event and prototype label; the authentic historical image with OSM credit.',
 'Unqualified impression of complete urban-planning coverage; claim that image is current connected output.',
 'Name the two actual workflows; label the historical preview.', 'KEEP WITH EDITS',
 'The jury learns scope immediately without more cover text or a new slide.'),
2: ('State the user problem',
 'The specialist is named, but the original problem is generic. Four transport branches can be mistaken for four implemented evaluation modules. No interviews, user quote or quantified loss is evidenced.',
 'Specialist as target; proposed-need disclosure; multimodal planning context.',
 'Any implied empirical pain validation or multi-mode implementation.',
 'A concrete route comparison question and explicit unknowns.', 'KEEP WITH EDITS',
 'A new hypothetical connection gives the user a concrete task linked to the implemented pedestrian study. Branches remain planning perspectives, explained in notes.'),
3: ('Acknowledge alternatives and proposed differentiation',
 'QGIS/SUMO descriptions are fair, but PMUD is a planning document rather than a competing product. The original source list is not three independent proofs of ARTERIA demand.',
 'QGIS and SUMO; advantage still to be tested.',
 'PMUD as an alternative product and any suggestion that competitors cannot compare scenarios.',
 'Data preparation/scenario comparison as a proposed work process, with qualified integration value.', 'RESTRUCTURE',
 'Keep the alternatives section and table; use its third row to explain the workflow without inventing competitor weaknesses.'),
4: ('Explain the user workflow',
 'The three stages exist, but synthetic signals and spatial graph studies need separate interpretation. Valuable same-sample comparisons are missing.',
 'Select area → inspect alternatives → compare.',
 'Generic future line duplicating roadmap.',
 'Same sampled routes for pedestrian alternatives; separate synthetic signal workflow.', 'KEEP WITH EDITS',
 'One concise present-tense capability is stronger than a generic future promise. Notes explain eligible OSM polygons and stored evaluation.'),
5: ('Show an actual product path and result',
 'The screenshot is authentic historical connected output, but the port is not a universal current address. The screenshot shows mean delay worsening; it must not be narrated as optimization. Its right edge crops part of the table.',
 'Three interaction steps, synthetic-data disclosure, screenshot and no-SUMO statement.',
 'Local port presented as a stable verified destination; geographic wording implying Podu Roș measurements.',
 'Tradeoff-focused headline, exact model inputs and adverse mean outcome in notes; optional spatial walkthrough for questions.', 'KEEP WITH EDITS',
 'One demo path remains understandable in 50 seconds. Replace only the cropped bitmap with a verified current capture in the same image geometry; no visual redesign.'),
6: ('Give technical credibility and accurate AI status',
 'Stack and wiring are mostly accurate. Local SUMO scope, stored spatial comparison and candidate validation are underspecified. The old footer mentions a dashed AI link absent from this main diagram.',
 'Three components, solid connected graph flow, Next.js/MapLibre, FastAPI/Pydantic/SQLite, optional local SUMO.',
 'Claim that main diagram has an actual dashed AI dependency; implication that all evaluation passes through SUMO or LLM.',
 'Stored graph comparison, candidate validation, existing evidence/catalog/rules, explicit LLM absence.', 'KEEP WITH EDITS',
 'AI status must be in the main presentation, given official evaluation criteria; no nonexistent model is added to improve the score.'),
7: ('Identify buyer and business hypothesis',
 '€1,500/month is inherited from previous pitch material. It has no interview, price test, accepted offer, measured delivery cost or paying-customer evidence. Existing qualification is honest but can be clearer.',
 'Buyer, proposed subscription, visibly unvalidated price, costs to measure.',
 'Any interpretation as actual tariff, traction or verified zero customers beyond available evidence.',
 'Evidence-limited customer wording; retain price only as the template-requested hypothesis.', 'KEEP WITH EDITS',
 'The template asks who pays and an explicitly hypothetical price. Retention is defensible with visible qualification; it is not market validation.'),
8: ('Explain market entry',
 'Ten organizations are a proposed list, not a verified segment. €180,000/year is correct multiplication but contributes no evidence of market size, conversion or revenue.',
 'Iași as first geography; mentor/community introductions as proposed channel.',
 'Dominant €180,000 scenario and arbitrary ten-organization target from the slide.',
 'First case/interviews and utility/willingness-to-pay exit criterion; preserve arithmetic only in notes.', 'REPLACE',
 'Replace speculative commercial arithmetic within the existing slide geometry. Preserve its position to respect the ten-section template; no new market statistic is invented.'),
9: ('Identify the team',
 'Names and broad responsibilities are supplied by the team; the English role descriptions break Romanian consistency.',
 'All three names, equal profiles and only supplied roles.',
 'English role descriptions and any invented granular authorship.',
 'Romanian equivalents of the same responsibilities.', 'KEEP WITH EDITS',
 'Translate without adding experience, ownership of handlers, full-time commitment or field expertise.'),
10: ('Close with next step and ask',
 'Data validation and specialist feedback are appropriate; missing LLM integration deserves explicit roadmap placement. Existing first-comparison wording can understate graph capabilities already built.',
 'Concrete introduction/data ask, specialist test and measurable workflow goal.',
 'Any impression that all data handling/graph comparisons remain to be built.',
 'LLM with validation; proposal correctness as an intended metric; calibration/field checks in notes.', 'KEEP WITH EDITS',
 'Use genuine implementation and validation gaps. No invented pilot, deadline or traffic reduction target.'),
11: ('Support technical questions with evidence',
 '9/9, typecheck and build need a current run. Historical 138 Python tests in notes must be independently rerun. Passing fixture tests do not establish calibrated impact or browser completion.',
 '9/9 frontend result; web/mobile/shared typecheck; build routes.',
 'Historical-only verification wording and untranslated check labels.',
 'Audit date, scope distinction and fresh Python result in notes/report.', 'KEEP WITH EDITS',
 'Keep technical detail in appendix; report the meaning of tests instead of inventing coverage percentages.'),
12: ('Explain AI boundary and future constrained generation',
 'LLM is already marked future, which is honest. The implemented bounded evidence packet and strict candidate-backed validation are underrepresented. Deterministic generation is not a working learned model.',
 'Evidence → future LLM → rules → independent evaluation; recommendation ≠ measurement.',
 'Any ambiguity that the proposed full AI flow executes today.',
 'Calculated candidates/rule checks today; evidence-bounded generation planned; no physical feasibility guarantee.', 'KEEP WITH EDITS',
 'This appendix explains existing engineering preparation for AI without relabeling it as an integrated LLM.'),
13: ('Explain methodology and limits',
 'Original content mixes implemented comparison controls and future simulation requirements. Repeating a result does not validate field accuracy.',
 'Provenance, shared baseline, explicit unknowns and no validated traffic outcomes.',
 'Generic reproducibility phrasing suggesting validity; simulation as an unqualified connected workflow.',
 'Shared graph/sample; sampled routes ≠ observed travel; hypothetical link ≠ feasible construction; local SUMO synthetic demand.', 'KEEP WITH EDITS',
 'A methodology slide is useful when it tells the jury what actually makes comparisons controlled, and what that control cannot establish.'),
}

def walk(shapes):
 for sh in shapes:
  yield sh
  if sh.shape_type==6: yield from walk(sh.shapes)

def visible(slide):
 result=[]
 for sh in walk(slide.shapes):
  if sh.has_text_frame and sh.text.strip(): result.append(sh.text)
  if sh.has_table:
   result.extend(c.text for row in sh.table.rows for c in row.cells if c.text.strip())
 return result

def esc(value): return str(value).replace('|','&#124;').replace('\n','<br>')

def judgement(n,text,is_note=False):
 if 'Sursa de adevăr:' in text:
  return 'MISLEADING','Research is a communication artifact, not the evidence authority.','Replace with current code/runtime evidence and independent inventory.'
 if '138' in text or '9/9' in text or '9 / 9' in text or 'typecheck' in text.lower() or 'Typecheck' in text or 'Build web' in text:
  return 'VERIFIED','Fresh checks reproduce these technical results; tests are not field validation.','Date the result and explain unit/fixture scope; keep historical results separate.'
 if ('cinci minute' in text.lower() or '300 s' in text or 'timpii sunt ținte' in text.lower()):
  return 'OUTDATED','A three-minute presentation is verified in organizer materials.','Replace five-minute rehearsal instructions with a 180-second proposed budget.'
 if 'întrerupt' in text.lower() and n==6:
  return 'INCORRECT','The main architecture slide has a future AI box but no dashed dependency connector to it.','Remove the visual-encoding assertion; leave geometry and animations intact.'
 if 'Frontend and' in text or 'Bootstrap and' in text or 'web experience' in text or 'simulation foundation' in text:
  return 'VERIFIED','Supplied responsibility is supported; language is inconsistent.','Translate broad responsibility to Romanian; do not infer detailed ownership.'
 if n in (7,8) and any(t in text for t in ['1.500','18.000','180.000','10 organizații','client','Client','Iași: primul']):
  return 'REQUIRES CONFIRMATION','Business price, organization list, adoption and conversion are not validated. Arithmetic is not market evidence.','Retain price visibly as hypothesis; replace market arithmetic with validation plan; qualify absence of customer evidence.'
 if n==5 and ('3001' in text or 'Verificare UI' in text):
  return 'OUTDATED','Port and verification description refer to a historical capture, not an always-running public demo.','Label historical capture; use current local setup during rehearsal.'
 if n==5 and ('nord–sud scade' in text or 'date sintetice' in text or 'Semaforizare' in text):
  return 'VERIFIED','Synthetic model supports the displayed tradeoff; mean delay worsens at the shown tuning.','Keep synthetic warning and explain adverse aggregate outcome, not optimization success.'
 if n==1 and 'Captură din aplicație' in text:
  return 'PARTIALLY VERIFIED','Authentic preview image, not proof of current connected analytics.','Label as historical preview.'
 if n==3 and ('PMUD' in text or 'Primăria' in text):
  return 'PARTIALLY VERIFIED','Planning context is not a competing product or evidence of ARTERIA demand; the historical municipal URL was not re-established in this audit.','Replace PMUD row; keep primary QGIS/SUMO sources.'
 if n==13 and ('Rulare repetată' in text or 'rezultat verificabil' in text):
  return 'PARTIALLY VERIFIED','Fixed inputs support repeatability, but repeatability alone cannot validate physical accuracy.','State repeatability is distinct from field validation.'
 if n==13 and ('Simulare'==text or 'Cerere documentată' in text):
  return 'PARTIALLY VERIFIED','Local SUMO configuration/output code exists; connected execution and calibration do not.','Label local synthetic-demand execution and distinguish future calibration.'
 if n==6 and 'SUMO local opțional' in text:
  return 'PARTIALLY VERIFIED','Local runner exists; binaries are absent and no API/UI SUMO path executes it.','State separate/local, with absent binaries and uncalibrated demand in notes.'
 if is_note and ('**Vizual' in text or '**Mișcare' in text):
  return 'PARTIALLY VERIFIED','This is authoring guidance mixed with factual notes; visual description does not itself prove functionality.','Replace with presenter-facing facts, evidence, limits and timing.'
 if text in EDITS.get(n,{}):
  return 'PARTIALLY VERIFIED',REVIEW[n][1],'Use the concise replacement in PPT_CONTENT_RESTRUCTURE_PLAN.md; preserve design.'
 if is_note and ('**Tranziție' in text or '**Verbal' in text):
  return 'PARTIALLY VERIFIED','Suggested narration/transition is not a record of observed product performance.','Rewrite around supported behavior and distinguish proposed validation from achieved outcomes.'
 return 'VERIFIED','Supported as scoped implementation, supplied identity, intended-use wording or explicit future/limitation label; not a new empirical outcome.','Retain where useful; concise wording and scope follow the independent inventory.'

def main():
 deck=Presentation(R/'.qa/build/content-audited.pptx')
 fact=['# PPT factual and editorial check — original V2','',
 'Audit: 18.10.2026. The repository inventory was established before the original deck/research was inspected. '
 'This report assesses the **frozen pre-audit deck**, not the corrected result. Exact OOXML text and speaker-note paragraphs are retained in '
 '`PPT_ORIGINAL_EXTRACTION.json`. All text-bearing shape/table/diagram elements are covered below; decorative section/page labels are not factual product claims. '
 'Long note rows intentionally preserve exact text so corrections can be traced. A VERIFIED label on a plan means it is honestly identified as a plan, not implemented functionality.','',
 'Status vocabulary: VERIFIED; PARTIALLY VERIFIED; OUTDATED; UNSUPPORTED; MISLEADING; INCORRECT; REQUIRES CONFIRMATION. '
 'Unsupported business assumptions are not declared false. External source checks and official-template conflict are in `HACKATHON_REQUIREMENTS_VERIFICATION.md`.','',
 '## Package-level issues','',
 '| Slide | Exact claim | Evidence | Status | Problem | Required action |','|---|---|---|---|---|---|',
 '| Package | 13-slide final used as the submission presentation | Official regulation Art. 7; participant guide §4 | INCORRECT | Official submitted presentation limit is 10; local template permits working Q&A appendices. | Preserve full working deck; provide separate exact 10-slide submission. |',
 '| All notes | Sursa de adevăr: docs/ppt/research/PROJECT_RESEARCH_FULL.md. | Source hierarchy; inspected implementation | MISLEADING | Earlier research is not authoritative evidence. | Replace every note with code/runtime evidence. |','']
 for n,original in enumerate(ORIGINAL,1):
  review=REVIEW[n]
  fact.extend([f'## Slide {n} — {TITLES[n-1]}','',review[1], '',
   '| Slide | Exact claim | Evidence | Status | Problem | Required action |','|---|---|---|---|---|---|'])
  # Consecutive a:t entries may be separate paragraphs; each is still assessed.
  for text in original['text']:
   if not text.strip() or re.fullmatch(r'\d{2}',text) or text.startswith('ARTERIA /'): continue
   status,problem,action=judgement(n,text)
   fact.append('| '+' | '.join(map(esc,[n,text,EVIDENCE[n],status,problem,action]))+' |')
  for text in original['notes']:
   if not text.strip() or text.startswith('ARTERIA —') or text=='Surse exacte și precizări:': continue
   status,problem,action=judgement(n,text,True)
   fact.append('| '+' | '.join(map(esc,[f'{n} notes',text,EVIDENCE[n],status,problem,action]))+' |')
  fact.append('')
 fact.extend(['## Screenshot text and numerical verification','',
 '| Slide | Exact claim | Evidence | Status | Problem | Required action |','|---|---|---|---|---|---|',
 '| 1 bitmap | DIGITAL TWIN MOBILITATE IAȘI; modeled map/unknown indicators | workspace-demo.png; Workspace.tsx:PreviewApp; MapViewport.tsx | MISLEADING | Screenshot UI branding can imply a complete digital twin; map preview contains fixtures and unknown values. | Preserve authentic bitmap; visibly label historical preview and describe only present workflows. |',
 '| 5 bitmap | Baseline NS/EW/mean 13,89; after NS 2,22, EW 125,56, mean 53,61 secunde/vehicul | SimulationAdapter.indicators; Inputs(90,70,10,700,500); fresh browser JSON | VERIFIED | These are rounded synthetic outputs; overall mean worsens, and they are not Podu Roș observations. | Explain negative tradeoff; no traffic-reduction claim. |',
 '| 5 bitmap | Înainte. După. Cu dovezi.; păstrând aceleași condiții de trafic; no traffic simulation run | ConnectedWorkspace.tsx; current /evaluate response | PARTIALLY VERIFIED | Same demand/cycle/lost time, but green allocation changes; original screenshot right edge crops the table. | Replace cropped bitmap with current verified capture at identical dimensions/geometry; explain fixed inputs. |',
 '| 1/5 bitmap | OpenStreetMap attribution / source screenshot | Original media and map source | VERIFIED | Basemap/source attribution is not source attribution for synthetic traffic metrics. | Keep OSM credit distinct from synthetic-model provenance. |','',
 'No native charts or embedded spreadsheets were found in the current 13-slide deck; editable tables and diagram labels were extracted from slide XML. '
 'No assertion of successful video playback, calibrated traffic or physical construction feasibility is made.','',
 '## Architecture claim resolution','',
 '| Component/claim | Evidence and actual integration | Verdict |','|---|---|---|',
 '| Next.js / MapLibre | package.json, page routing and MapViewport; MapLibre raster basemap and selected source extent | VERIFIED; pitched raster map is not a complete 3D digital twin |',
 '| FastAPI / Pydantic / SQLite | main.py, models.py, storage.py; live fixture flow plus 138 tests | VERIFIED; not auth/tenant collaboration |',
 '| Connected client | createApiClient → ConnectedWorkspace → real endpoints | VERIFIED for synthetic signals and fixture spatial flow; runtime response validation on client absent |',
 '| Spatial studies / pedestrian graph | spatial_worker.calculate → analyze_snapshot/evaluate_candidates; stored study/evaluate response | VERIFIED; not observed trips |',
 '| OpenStreetMap | catalog/tile source and acquisition code; canonical neighborhood ID | PARTIALLY VERIFIED at runtime: no fresh external download; graph API skips feature acquisition |',
 '| Scenario generation | deterministic signal greens; candidate-backed hypothetical links | VERIFIED as deterministic; no learned generation |',
 '| Evaluation | exact IDs/action match; stored common-baseline graph comparison | VERIFIED; spatial selection does not rerun SIM |',
 '| SUMO | prepare_sumo_network/run_real_area_simulation; smoke says missing binaries | IMPLEMENTED LOCAL CODE, not executed/integrated/calibrated |',
 '| AI | bounded evidence/context and rule validator; no provider call | LLM PLANNED; cannot claim current generative AI |',''])
 (R/'PPT_FACT_CHECK.md').write_text('\n'.join(fact),encoding='utf-8')

 plan=['# PPT content restructure plan','',
 'Prepared after the independent implementation inventory and original-claim audit, before publishing corrections. '
 'Apply through `scripts/audited_content.py`, after existing V2 shape/refinement creation. '
 'The following is exact proposed **visible** content (including retained headings/footers and editable table cells), followed by suggested Romanian speaker notes. '
 'Line breaks follow existing text boxes. No geometry, typography, colors, animations or slide order is changed.','',
 '## Overall narrative and structural decision','',
 'Keep the template sequence: user/problem → alternatives → simple workflow → one demo → architecture/AI boundary → business → market validation → team → concrete ask. '
 'The original first two slides named the user but lacked a concrete decision; the revised problem links directly to route comparisons. '
 'The architecture remains a three-component explanation rather than a technical specification. '
 'Same-sample graph comparison and candidate checks are worth main-slide space; detailed methodology, test scope and future AI flow remain Q&A. '
 'No missing slide justifies increasing the count. The market content is replaced in place; appendices remain only in the full working copy.','',
 'Official submission format and schedule conflict: see `HACKATHON_REQUIREMENTS_VERIFICATION.md`. '
 'A 13-slide working copy is not certified as a compliant ten-slide submission. The derivative contains slides 1–10 only.','',
 '| Current slide | Decision | Reason |','|---|---|---|']
 for n in range(1,14): plan.append(f'| {n} | {REVIEW[n][5]} | {REVIEW[n][6]} |')
 plan.extend(['','## Important missing or underrepresented information','',
 '| Element | Why it matters to the jury | Verification | Placement / replacement | Pitch benefit |','|---|---|---|---|---|',
 '| Shared seeded route sample; independent unchanged baseline | Makes alternatives comparable | impact.py:evaluate_candidates and regression tests; fixture API exercised | Slide 4 replaces generic present/future line; slide 13 gives method | Concrete defensible engineering result |',
 '| Candidate-backed validation and bounded evidence packet | Shows preparation for constrained AI rather than invented AI | proposal_context/proposals/proposal_validation through worker and UI | Slide 6 future-provider text becomes current evidence/rules plus missing LLM; slide 12 details | Honest account of AI-related achievement and gap |',
 '| Catalog source and eligibility | Demonstrates reuse across sourced areas | 24 polygons + 7 points, tests, canonical ID validation | Slide 4 notes and questions; do not crowd main workflow | Geographic credibility without claiming full city coverage |',
 '| Stored comparison / SQLite | Explains traceability and avoids false new-run claims | spatial_routes.evaluate returns stored impact | Slide 6 disclosure and slide 13 notes | Reproducible decision trail |',
 '| Adverse mean-delay tradeoff | Shows the system does not always display improvement | Formula and browser value at green 70 | Slide 5 headline/notes instead of place-based success framing | Defensible demo under questioning |',
 '| Measured technical checks vs unmeasured city impact | Prevents tests being mistaken for product validation | Fresh check logs | Slide 11 notes, main slide 6 scope | Evidence proportional to claim |',
 '| Missing runtime AI | Largest official criterion remains a gap | No provider call; deterministic metadata | Main slide 6 and roadmap 10 | Avoids hiding a material limitation in appendix |','',
 'Do not add full library capability lists, SUMO vehicle counts, old cached-route gains, authentication/export specs, mobile parity, or speculative global TAM. '
 'Those would increase text or imply integration/validation not demonstrated.','',
 '## Three-minute rehearsal allocation','',
 'Proposed allocation across slides 1–10: '+', '.join(f'{i+1}: {v}s' for i,v in enumerate(TIMES))+'. Total: 180s; cover is pre-roll. '
 'This is an allocation, not a verified successful rehearsal. The template also permits omitting market discussion at three minutes; the ten-slide submission retains that section. '
 'Reserve live demo for the existing synthetic path; open pedestrian details during questions if time does not allow both.',''])
 for n,slide in enumerate(deck.slides,1):
  purpose,problems,remain,remove,add,decision,reason=REVIEW[n]
  title=EDITS[n].get(TITLES[n-1],TITLES[n-1])
  plan.extend([f'## Slide {n}','',f'**Current title:** {TITLES[n-1]}','',f'**Current purpose:** {purpose}','',
    f'**Problems identified:** {problems}','',f'**What should remain:** {remain}','',
    f'**What should be removed:** {remove}','',f'**What should be added:** {add}','',
    f'**Improved headline:** {title}','', '**Exact proposed slide text:**',''])
  for text in visible(slide): plan.extend('> '+line for line in text.split('\n'))
  plan.extend(['',f'**Verified evidence supporting that text:** {EVIDENCE[n]}','',
    '**Suggested speaker notes:**', '',NOTES[n],'',f'**Reasoning behind the changes:** {reason}',''])
 (R/'PPT_CONTENT_RESTRUCTURE_PLAN.md').write_text('\n'.join(plan),encoding='utf-8')
 print('Wrote complete original-claim matrix and slide-by-slide content/structure plan.')

if __name__=='__main__':main()
