"""Compare against the frozen user deck; content edits may not alter design/motion."""
import argparse
import hashlib
import json
from pathlib import Path
import zipfile
from lxml import etree as E
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}

def digest(blob):return hashlib.sha256(blob).hexdigest()
def canonical(blob,blank_text=False):
 root=E.fromstring(blob)
 if blank_text:
  for t in root.xpath('//a:t',namespaces=NS):t.text=''
 return E.tostring(root,method='c14n')

def verify(original,corrected,submission):
 result={'backup_sha256':digest(original.read_bytes()),'corrected_sha256':digest(corrected.read_bytes()),'slides':[]}
 with zipfile.ZipFile(original) as before,zipfile.ZipFile(corrected) as after:
  assert set(before.namelist())==set(after.namelist())
  changed=[]
  for name in before.namelist():
   b,a=before.read(name),after.read(name)
   if b==a:continue
   changed.append(name)
   if name.startswith('ppt/media/'):
    assert name=='ppt/media/image2.jpg',f'Unexpected media correction: {name}'
    assert a==(ROOT/'assets/connected-comparison-audited.jpg').read_bytes()
    continue
   assert name.startswith(('ppt/slides/slide','ppt/notesSlides/notesSlide')) and name.endswith('.xml'),name
   if name.startswith('ppt/slides/'):
    assert canonical(b,True)==canonical(a,True),f'Non-text change: {name}'
  for i in range(1,14):
   name=f'ppt/slides/slide{i}.xml'; b=E.fromstring(before.read(name));a=E.fromstring(after.read(name))
   for tag in ('timing','transition'):
    bx=b.find('.//p:'+tag,NS);ax=a.find('.//p:'+tag,NS)
    assert bx is not None and ax is not None
    assert E.tostring(bx,method='c14n')==E.tostring(ax,method='c14n')
   result['slides'].append({'slide':i,'nontext_xml_identical':True,'timing_identical':True,'transition_identical':True})
  result['changed_parts']=changed
  result['unchanged_theme_layout_geometry_styles']=True
  result['media_exception']='Slide 5 comparison bitmap only: real current uncropped capture; original picture geometry unchanged.'
 with zipfile.ZipFile(corrected) as full,zipfile.ZipFile(submission) as sub:
  for i in range(1,11):
   name=f'ppt/slides/slide{i}.xml';assert full.read(name)==sub.read(name)
  result['submission_main_slides_byte_identical']=True
 assert len(Presentation(corrected).slides)==13
 assert len(Presentation(submission).slides)==10
 p=Presentation(corrected)
 for slide in p.slides:assert slide.notes_slide.notes_text_frame.text.strip()
 result['notes_complete']=True;result['full_slide_count']=13;result['submission_slide_count']=10;result['passed']=True
 return result

if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--original',type=Path,default=ROOT/'archive/v2-before-content-audit-20261018/Template-Pitch-Hackathon.v2.pptx')
 ap.add_argument('--deck',type=Path,default=ROOT/'reports/.qa/build/content-audited.pptx')
 ap.add_argument('--submission',type=Path,default=ROOT/'reports/.qa/build/ARTERIA-Submission.pptx')
 args=ap.parse_args();r=verify(args.original,args.deck,args.submission)
 (ROOT/'reports/content-preservation.json').write_text(json.dumps(r,indent=2),encoding='utf-8')
 print(json.dumps({k:v for k,v in r.items() if k not in ('slides','changed_parts')},indent=2))
