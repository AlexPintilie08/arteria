"""Artifact QA: native object inventory, text bounds, source preservation, previews."""
from pathlib import Path
import json, hashlib, re, zipfile
from pptx import Presentation
from PIL import Image, ImageDraw

ROOT=Path(__file__).resolve().parent
p=Presentation(ROOT/'Template-Pitch-Hackathon.pptx')
checks=json.loads((ROOT/'preview/text-bounds.json').read_text(encoding='utf-8-sig'))
overflow=[x for x in checks if x['boundHeight']>x['height']+1 or x['boundWidth']>x['width']+1]
out=[];texts=[];counts={'textboxes':0,'tables':0,'pictures':0,'connectors':0,'notes':0}
for i,s in enumerate(p.slides,1):
    counts['notes']+=bool(s.notes_slide.notes_text_frame.text.strip())
    for sh in s.shapes:
        if sh.has_text_frame:counts['textboxes']+=1;texts.append(sh.text)
        if sh.has_table:
            counts['tables']+=1
            texts.extend(c.text for row in sh.table.rows for c in row.cells)
        counts['pictures']+=sh.shape_type==13
        counts['connectors']+=sh.shape_type==9
        if sh.left<0 or sh.top<0 or sh.left+sh.width>p.slide_width+10 or sh.top+sh.height>p.slide_height+10:out.append({'slide':i,'name':sh.name})
alltext='\n'.join(texts)
residuals=re.findall(r'\[NUME|\[CINE|GHIDAJ — NU INTRĂ|RĂMÂNE PE SLIDE|lorem ipsum',alltext,re.I)
files=sorted((ROOT/'preview').glob('Slide*.PNG'),key=lambda f:int(f.stem[5:]))
for j in range(0,len(files),4):
    sheet=Image.new('RGB',(1600,940),'white')
    for k,f in enumerate(files[j:j+4]):
        im=Image.open(f);assert im.size==(1920,1080);im.thumbnail((800,450))
        x=k%2*800;y=k//2*470;sheet.paste(im,(x,y));ImageDraw.Draw(sheet).text((x+10,y+450),f.stem,fill='black')
    sheet.save(ROOT/'preview'/f'contact-{j//4+1}.jpg')
original=json.loads((ROOT/'Template-Pitch-Hackathon.backup.inspection.json').read_text(encoding='utf-8'))
backup_hash=hashlib.sha256((ROOT/'Template-Pitch-Hackathon.backup.pptx').read_bytes()).hexdigest()
with zipfile.ZipFile(ROOT/'Template-Pitch-Hackathon.pptx') as z:assert z.testzip() is None
report={'slides':len(p.slides),'slide_size_inches':[p.slide_width/914400,p.slide_height/914400],
    'objects':counts,'text_bounds_checked':len(checks),'text_overflow':overflow,'out_of_canvas':out,
    'remaining_template_prompts':residuals,'rendered_slides':len(files),'render_size':[1920,1080],
    'backup_sha256':backup_hash,'backup_unchanged':backup_hash==original['sha256'],
    'pptx_sha256':hashlib.sha256((ROOT/'Template-Pitch-Hackathon.pptx').read_bytes()).hexdigest()}
(ROOT/'qa-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
assert len(p.slides)==13 and len(files)==13 and not overflow and not out and not residuals
assert report['backup_unchanged'] and counts['notes']==13 and counts['tables']==2 and counts['pictures']==2
print(json.dumps(report,ensure_ascii=False,indent=2))
