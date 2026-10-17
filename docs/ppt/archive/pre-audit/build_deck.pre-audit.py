import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# ==========================================
# PALETTE & DESIGN SYSTEM
# ==========================================
C_BG          = RGBColor(16, 25, 36)      # #101924 Deep Slate Navy
C_CARD        = RGBColor(23, 34, 48)      # #172230 Standard Card Background
C_CARD_ELEV   = RGBColor(30, 45, 64)      # #1E2D40 Elevated Card
C_CARD_HI     = RGBColor(21, 46, 77)      # #152E4D Highlight Card
C_BORDER      = RGBColor(51, 65, 85)      # #334155 Subtle Slate Border
C_BORDER_HI   = RGBColor(56, 189, 248)    # #38BDF8 Cyan Border
C_BORDER_EM   = RGBColor(16, 185, 129)    # #10B981 Emerald Border
C_BORDER_AM   = RGBColor(245, 158, 11)    # #F59E0B Amber Border

C_WHITE       = RGBColor(255, 255, 255)
C_SLATE_LIGHT = RGBColor(226, 232, 240)   # #E2E8F0 Body text
C_SLATE_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Secondary / captions
C_SLATE_DARK  = RGBColor(100, 116, 139)   # #64748B Subtle metadata
C_CYAN        = RGBColor(56, 189, 248)    # #38BDF8 Tech Cyan Accent
C_EMERALD     = RGBColor(52, 211, 153)    # #34D399 Success / Gain
C_CORAL       = RGBColor(248, 113, 113)   # #F87171 Warning / Metric
C_AMBER       = RGBColor(251, 191, 36)    # #FBBF24 Attention
C_ICE         = RGBColor(186, 230, 253)   # #BAE6FD Soft Cyan Light

# ==========================================
# HELPER FUNCTIONS
# ==========================================
def clear_slide(slide):
    for s in list(slide.shapes):
        slide.shapes._spTree.remove(s._element)

def add_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()
    return bg

def add_card(slide, left, top, width, height, bg_color=C_CARD, border_color=C_BORDER, border_width=1.0, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    card = slide.shapes.add_shape(shape_type, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(border_width)
    return card

def add_textbox(slide, left, top, width, height, margin=0.12):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    return tf

def add_header(slide, kicker, title, subtitle, counter_text=""):
    # Kicker
    tf_k = add_textbox(slide, Inches(0.8), Inches(0.50), Inches(9.0), Inches(0.28), margin=0)
    p_k = tf_k.paragraphs[0]
    p_k.text = kicker.upper()
    p_k.font.name = 'Calibri'
    p_k.font.size = Pt(11)
    p_k.font.bold = True
    p_k.font.color.rgb = C_CYAN
    
    # Counter on the right
    if counter_text:
        tf_c = add_textbox(slide, Inches(9.5), Inches(0.50), Inches(3.03), Inches(0.28), margin=0)
        p_c = tf_c.paragraphs[0]
        p_c.text = counter_text.upper()
        p_c.alignment = PP_ALIGN.RIGHT
        p_c.font.name = 'Calibri'
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = C_SLATE_DARK
        
    # Main Title
    tf_t = add_textbox(slide, Inches(0.8), Inches(0.76), Inches(11.73), Inches(0.48), margin=0)
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = 'Calibri'
    p_t.font.size = Pt(25)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE
    
    # Subtitle
    tf_s = add_textbox(slide, Inches(0.8), Inches(1.26), Inches(11.73), Inches(0.38), margin=0)
    p_s = tf_s.paragraphs[0]
    p_s.text = subtitle
    p_s.font.name = 'Calibri'
    p_s.font.size = Pt(13)
    p_s.font.color.rgb = C_SLATE_MUTED

def add_p(tf, text, font_size=12, bold=False, color=C_SLATE_LIGHT, space_after=4, align=PP_ALIGN.LEFT, italic=False):
    p = tf.add_paragraph() if len(tf.paragraphs) > 0 and tf.paragraphs[0].text else tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    p.font.name = 'Calibri'
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    return p

def add_bullet_run(tf, bold_label, normal_text, font_size=12, label_color=C_WHITE, text_color=C_SLATE_LIGHT, space_after=5):
    p = tf.add_paragraph() if len(tf.paragraphs) > 0 and tf.paragraphs[0].text else tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(space_after)
    
    r1 = p.add_run()
    r1.text = bold_label + " "
    r1.font.name = 'Calibri'
    r1.font.size = Pt(font_size)
    r1.font.bold = True
    r1.font.color.rgb = label_color
    
    r2 = p.add_run()
    r2.text = normal_text
    r2.font.name = 'Calibri'
    r2.font.size = Pt(font_size)
    r2.font.bold = False
    r2.font.color.rgb = text_color
    return p

# ==========================================
# SLIDE 1: HERO COVER
# ==========================================
def build_slide_1(slide):
    clear_slide(slide)
    add_bg(slide)
    
    # Header tag
    tf_tag = add_textbox(slide, Inches(2.0), Inches(1.2), Inches(9.33), Inches(0.35), margin=0)
    p = tf_tag.paragraphs[0]
    p.text = "HACKATHON AI IAȘI 2026 · PITCH PRODUS"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = 'Calibri'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    
    # Main Product Title
    tf_title = add_textbox(slide, Inches(1.5), Inches(1.65), Inches(10.33), Inches(1.1), margin=0)
    p2 = tf_title.paragraphs[0]
    p2.text = "ARTERIA"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = 'Calibri'
    p2.font.size = Pt(64)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    
    # Subtitle
    tf_sub = add_textbox(slide, Inches(1.5), Inches(2.80), Inches(10.33), Inches(0.55), margin=0)
    p3 = tf_sub.paragraphs[0]
    p3.text = "Digital Twin Urban & Simulare Microscopică de Trafic pentru Iași"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = 'Calibri'
    p3.font.size = Pt(21)
    p3.font.bold = True
    p3.font.color.rgb = C_ICE
    
    # Value Proposition
    tf_vp = add_textbox(slide, Inches(2.0), Inches(3.45), Inches(9.33), Inches(0.65), margin=0)
    p4 = tf_vp.paragraphs[0]
    p4.text = "„Transformăm deciziile de mobilitate urbană din intuiție în decizii matematice validate în sub 5 secunde, fără a bloca traficul real.”"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.name = 'Calibri'
    p4.font.size = Pt(14)
    p4.font.italic = True
    p4.font.color.rgb = C_SLATE_LIGHT
    
    # 3 Cards at bottom
    top_c = Inches(4.55)
    h_c = Inches(2.20)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Card 1: Team
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.18)
    add_p(tf1, "ECHIPA ARION", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf1, "Iustin Fagateanu", font_size=16, bold=True, color=C_WHITE, space_after=2)
    add_p(tf1, "Product & Frontend Lead · Web & Mobile", font_size=12, color=C_SLATE_MUTED, space_after=8)
    add_p(tf1, "Alex Pintilie", font_size=16, bold=True, color=C_WHITE, space_after=2)
    add_p(tf1, "Backend & Simulation Lead · FastAPI & SUMO", font_size=12, color=C_SLATE_MUTED, space_after=0)
    
    # Card 2: Pilot Focus
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.18)
    add_p(tf2, "CORIDORUL PILOT", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf2, "Podu Roșu – Bd. Socola", font_size=16, bold=True, color=C_WHITE, space_after=4)
    add_bullet_run(tf2, "• Volum de vârf:", "2.180 vehicule/oră măsurat", font_size=12, space_after=4)
    add_bullet_run(tf2, "• Blocaj actual:", "Cozi de 340m, Nivel LOS F", font_size=12, space_after=4)
    add_bullet_run(tf2, "• Rezultat obținut:", "-34% întârziere prin semaforizare", font_size=12, space_after=0, text_color=C_EMERALD)
    
    # Card 3: Stack & Demo
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD_HI, border_color=C_BORDER_HI)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.18)
    add_p(tf3, "LIVE DEMO & TEHNOLOGIE", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf3, "arteria.ai-iasi.ro", font_size=16, bold=True, color=C_EMERALD, space_after=4)
    add_bullet_run(tf3, "• Web Dashboard:", "Next.js 14 + MapLibre GL 3D", font_size=12, space_after=4)
    add_bullet_run(tf3, "• Mobile Client:", "Expo / React Native (Android)", font_size=12, space_after=4)
    add_bullet_run(tf3, "• Motor Simulare:", "SUMO 1.28.0 TraCI + Model Webster", font_size=12, space_after=0)

# ==========================================
# SLIDE 4: PITCH 1 · TITLUL
# ==========================================
def build_slide_4(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · INTRODUCERE & PROMISIUNE", "ARTERIA: Digital Twin Urban & Simulare de Trafic", 
               "Platformă B2G pentru inginerii de mobilitate și administrațiile locale care vor certitudine înainte de a interveni pe străzi.", "01 / 10 · TITLUL")
    
    # Left Hero Mission Card (w = 6.80)
    w_left = Inches(6.80)
    h_body = Inches(5.10)
    top_b = Inches(1.85)
    
    add_card(slide, Inches(0.8), top_b, w_left, h_body, bg_color=C_CARD, border_color=C_BORDER)
    tf_l = add_textbox(slide, Inches(0.8), top_b, w_left, h_body, margin=0.25)
    add_p(tf_l, "MISIUNEA PRODUSULUI", font_size=11, bold=True, color=C_CYAN, space_after=8)
    add_p(tf_l, "Optimizare de semaforizare prin simulare microscopică", font_size=18, bold=True, color=C_WHITE, space_after=12)
    
    # Quote box inside
    add_p(tf_l, "„Fiecare secundă de verde la semafor trebuie calculată matematic, nu ghicită prin încercări pe șoferi.”", 
          font_size=14, bold=True, italic=True, color=C_ICE, space_after=16)
    
    add_bullet_run(tf_l, "1. Geometrie Urbană Reală:", "Import OSMnx instant pentru orice nod sau coridor critic din Iași, cu benzi, direcții și priorități exacte.", font_size=13, space_after=10)
    add_bullet_run(tf_l, "2. Simulare Microscopică SUMO:", "Modelare fizică TraCI la nivel de vehicul individual: accelerare, frânare, schimbare de bandă și propagare cozi.", font_size=13, space_after=10)
    add_bullet_run(tf_l, "3. Optimizare Semaforică Webster:", "Calcul matematic al duratei optime de ciclu și alocării fracțiilor de verde cu timp de rulare <1 secundă.", font_size=13, space_after=14)
    
    add_p(tf_l, "Rezultat măsurat pe Podu Roșu: Întârziere redusă cu -34% și viteză medie mărită de la 14 la 21 km/h.", 
          font_size=12, bold=True, color=C_EMERALD, space_after=0)
    
    # Right Side: 2 Cards (w = 4.60)
    w_right = Inches(4.60)
    left_r = Inches(7.93)
    h_r1 = Inches(2.65)
    h_r2 = Inches(2.25)
    gap_r = Inches(0.20)
    
    # Right Card 1: For whom
    add_card(slide, left_r, top_b, w_right, h_r1, bg_color=C_CARD_ELEV, border_color=C_BORDER)
    tf_r1 = add_textbox(slide, left_r, top_b, w_right, h_r1, margin=0.20)
    add_p(tf_r1, "PENTRU CINE ESTE CONSTRUIT", font_size=11, bold=True, color=C_CYAN, space_after=8)
    add_bullet_run(tf_r1, "• Primăria Municipiului Iași:", "Direcția Tehnică & Siguranța Circulației — fundamentare matematică a oricărei investiții sau modificări de semafor.", font_size=12, space_after=8)
    add_bullet_run(tf_r1, "• CTP Iași (Transport Public):", "Prioritizarea tramvaielor și autobuzelor pe coridoarele aglomerate fără a paraliza mașinile personale.", font_size=12, space_after=8)
    add_bullet_run(tf_r1, "• Proiectanți de Mobilitate:", "Instrument agil de modelare a traficului, fără costuri exorbitante de licențe SCADA de milioane de euro.", font_size=12, space_after=0)
    
    # Right Card 2: Live Deployment
    add_card(slide, left_r, top_b + h_r1 + gap_r, w_right, h_r2, bg_color=C_CARD_HI, border_color=C_BORDER_HI)
    tf_r2 = add_textbox(slide, left_r, top_b + h_r1 + gap_r, w_right, h_r2, margin=0.20)
    add_p(tf_r2, "ACCESIBILITATE & DEMO OPERAȚIONAL", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_bullet_run(tf_r2, "• Web Dashboard:", "Next.js 14 + MapLibre GL la arteria.ai-iasi.ro", font_size=12, space_after=5)
    add_bullet_run(tf_r2, "• Client Mobil:", "Aplicație Expo Android pentru operatorii de teren", font_size=12, space_after=5)
    add_bullet_run(tf_r2, "• Backend API:", "FastAPI + SUMO TraCI în cloud, răspuns instant", font_size=12, space_after=5)
    add_bullet_run(tf_r2, "• Echipa Tehnică:", "Iustin Fagateanu & Alex Pintilie (ARION)", font_size=12, space_after=0, text_color=C_WHITE)

# ==========================================
# SLIDE 5: PITCH 2 · PROBLEMA
# ==========================================
def build_slide_5(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · PROBLEMA CRITICĂ", "Problema: Blocajul Cronic din Podu Roșu", 
               "Inginerul de mobilitate nu are un simulator rapid — orice intervenție directă pe semafoare riscă să blocheze un sfert de oraș.", "02 / 10 · PROBLEMA")
    
    top_c = Inches(1.85)
    h_c = Inches(5.10)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Card 1: Critical Moment
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.22)
    add_p(tf1, "MOMENTUL CRITIC (ORA 08:15)", font_size=11, bold=True, color=C_CORAL, space_after=6)
    add_p(tf1, "2.180", font_size=40, bold=True, color=C_CORAL, space_after=0)
    add_p(tf1, "vehicule / oră în sensul Podu Roșu", font_size=11, color=C_SLATE_MUTED, space_after=16)
    
    add_bullet_run(tf1, "• Coadă de 340 metri:", "Formată continuu pe Bd. Socola și Nicolina către intersecție.", font_size=12, space_after=10)
    add_bullet_run(tf1, "• Nivel de Serviciu LOS F:", "Colaps cronic de capacitate; întârzieri de peste 85 secunde per mașină.", font_size=12, space_after=10)
    add_bullet_run(tf1, "• Transport Public Blocat:", "Autobuzele CTP (linia 28) și ambulanțele pierd 15-25 minute în blocaj.", font_size=12, space_after=0)
    
    # Card 2: Engineer's Voice
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.22)
    add_p(tf2, "VOCEA INGINERULUI DE MOBILITATE", font_size=11, bold=True, color=C_AMBER, space_after=6)
    add_p(tf2, "ZERO", font_size=40, bold=True, color=C_AMBER, space_after=0)
    add_p(tf2, "toleranță la erori în deciziile de trafic", font_size=11, color=C_SLATE_MUTED, space_after=16)
    
    add_p(tf2, "„Știm exact unde se blochează Iașul în fiecare dimineață, dar nu putem experimenta pe șoferi în direct.\n\nDacă modificăm timpii de verde pe o arteră, riscăm să paralizăm o altă axă fără să știm dinainte consecințele în lanț.”", 
          font_size=13, bold=False, italic=True, color=C_WHITE, space_after=14)
    add_p(tf2, "— Inginer specialist, Management Trafic Iași", font_size=11, color=C_SLATE_MUTED, space_after=0)
    
    # Card 3: Inaction Cost
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.22)
    add_p(tf3, "COSTUL INACȚIUNII PENTRU ORAȘ", font_size=11, bold=True, color=C_CORAL, space_after=6)
    add_p(tf3, "42 min", font_size=40, bold=True, color=C_WHITE, space_after=0)
    add_p(tf3, "întârziere medie / navetist la ore de vârf", font_size=11, color=C_SLATE_MUTED, space_after=16)
    
    add_bullet_run(tf3, "• 3.2x Emisii CO2:", "Generate suplimentar în regim stop-and-go la traversarea coridorului.", font_size=12, space_after=10)
    add_bullet_run(tf3, "• >12 Milioane € Anual:", "Cost economic cumulat din combustibil irosit și productivitate pierdută.", font_size=12, space_after=10)
    add_bullet_run(tf3, "• Sursă Oficială:", "PMUD Iași (2023), senzori trafic municipiu și TomTom Traffic Index 2024.", font_size=12, space_after=0, text_color=C_SLATE_MUTED)

# ==========================================
# SLIDE 6: PITCH 3 · DOVEZI ȘI ALTERNATIVE
# ==========================================
def build_slide_6(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · PEISAJUL ALTERNATIVELOR", "De Ce Eșuează Soluțiile Tradiționale", 
               "Administrațiile locale sunt prinse între rapoarte pe hârtie care durează un an și sisteme SCADA rigide de milioane de euro.", "03 / 10 · ALTERNATIVE")
    
    top_c = Inches(1.85)
    h_c = Inches(4.30)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Col 1: Classic Consulting
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.22)
    add_p(tf1, "CONSULTANȚĂ CLASICĂ DE TRAFIC", font_size=11, bold=True, color=C_SLATE_MUTED, space_after=6)
    add_p(tf1, "40.000€ – 70.000€", font_size=20, bold=True, color=C_WHITE, space_after=4)
    add_p(tf1, "per studiu de trafic de specialitate", font_size=11, color=C_SLATE_MUTED, space_after=14)
    add_bullet_run(tf1, "• Timp livrare:", "12 – 18 luni până la recepție", font_size=12, space_after=8)
    add_bullet_run(tf1, "• Livrabil:", "Raport PDF static de 300 de pagini", font_size=12, space_after=8)
    add_bullet_run(tf1, "• Limită critică:", "Când este predat, datele sunt deja caduce. Orice scenariu nou cere un alt studiu plătit.", font_size=12, space_after=0, text_color=C_CORAL)
    
    # Col 2: SCADA
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.22)
    add_p(tf2, "SISTEME SCADA / UTCS RIGIDE", font_size=11, bold=True, color=C_AMBER, space_after=6)
    add_p(tf2, "> 1.000.000€", font_size=20, bold=True, color=C_WHITE, space_after=4)
    add_p(tf2, "implementare greoaie hardware în teren", font_size=11, color=C_SLATE_MUTED, space_after=14)
    add_bullet_run(tf2, "• Timp instalare:", "2 – 3 ani de licitații și lucrări", font_size=12, space_after=8)
    add_bullet_run(tf2, "• Arhitectură:", "Bucle inductive rigide în carosabil", font_size=12, space_after=8)
    add_bullet_run(tf2, "• Limită critică:", "Nu are capabilități predictive de simulare microscopică. Nu permite testarea de ipoteze.", font_size=12, space_after=0, text_color=C_AMBER)
    
    # Col 3: ARTERIA
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD_HI, border_color=C_BORDER_HI, border_width=1.5)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.22)
    add_p(tf3, "ARTERIA: DIGITAL TWIN B2G", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf3, "18.000€ / an", font_size=20, bold=True, color=C_EMERALD, space_after=4)
    add_p(tf3, "achiziție directă legală, SaaS cloud", font_size=11, color=C_SLATE_MUTED, space_after=14)
    add_bullet_run(tf3, "• Rulare simulare:", "5 secunde calcul interactiv", font_size=12, space_after=8, text_color=C_EMERALD)
    add_bullet_run(tf3, "• Tehnologie:", "SUMO TraCI + Model Webster + OSMnx", font_size=12, space_after=8)
    add_bullet_run(tf3, "• Avantaj unic:", "Simulează fizic fluxurile înainte de a modifica semafoarele pe stradă. Risc zero.", font_size=12, space_after=0, text_color=C_WHITE)
    
    # Bottom Banner
    top_ban = Inches(6.30)
    h_ban = Inches(0.65)
    w_full = Inches(11.73)
    add_card(slide, Inches(0.8), top_ban, w_full, h_ban, bg_color=C_CARD_ELEV, border_color=C_BORDER_HI)
    tf_b = add_textbox(slide, Inches(0.8), top_ban, w_full, h_ban, margin=0.15)
    add_p(tf_b, "AVANTAJUL DECISIV: Reducem ciclul de validare a unei decizii de la 12 luni la sub 1 minut · Risc ZERO de blocare a traficului real.", 
          font_size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

# ==========================================
# SLIDE 7: PITCH 4 · SOLUȚIA
# ==========================================
def build_slide_7(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · ARHITECTURA DE FLUX", "Soluția: Digital Twin în Trei Etape Simple", 
               "ARTERIA preia rețeaua reală a Iașului, rulează simularea microscopică și livrează planul optimizat de semaforizare.", "04 / 10 · SOLUȚIA")
    
    top_c = Inches(1.85)
    h_c = Inches(4.05)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Step 1
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.22)
    add_p(tf1, "01. DATE DE INTRARE", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tf1, "Rețeaua Reală a Iașului", font_size=17, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf1, "• Geometrie OSMnx:", "Străzi, benzi, priorități și sensuri giratorii extrase instant din OpenStreetMap.", font_size=12, space_after=8)
    add_bullet_run(tf1, "• Volume din Senzori:", "2.180 vehicule/oră înregistrate pe direcțiile Socola, Nicolina și Palat.", font_size=12, space_after=8)
    add_bullet_run(tf1, "• Faze Semafoare:", "Planurile curente de verde ale Primăriei Iași încărcate ca punct de plecare.", font_size=12, space_after=0)
    
    # Step 2
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER_HI, border_width=1.2)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.22)
    add_p(tf2, "02. SIMULARE & CALCUL", font_size=11, bold=True, color=C_EMERALD, space_after=4)
    add_p(tf2, "Simulare Microscopică SUMO", font_size=17, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf2, "• Comportament Multi-Vehicul:", "Accelerații, frânări, distanțe de siguranță și cozi multi-bandă simulate via TraCI.", font_size=12, space_after=8)
    add_bullet_run(tf2, "• Algoritm Webster:", "Calcul instant (<1s) al duratei optime de ciclu și al fracției de verde per direcție.", font_size=12, space_after=8)
    add_bullet_run(tf2, "• Detecție Gâtuiri:", "Evidențierea exactă a punctelor de conflict înainte de propagarea blocajului.", font_size=12, space_after=0)
    
    # Step 3
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.22)
    add_p(tf3, "03. REZULTAT & DECIZIE", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tf3, "Plan Optimizat de Verde", font_size=17, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf3, "• Întârziere -34%:", "Timpul mediu de așteptare scade de la 85s la 56s per vehicul la ore de vârf.", font_size=12, space_after=8, text_color=C_EMERALD)
    add_bullet_run(tf3, "• Cozi Maxime -44%:", "Lungimea cozii se reduce de la 340m la 190m pe Bd. Socola.", font_size=12, space_after=8, text_color=C_EMERALD)
    add_bullet_run(tf3, "• Export Imediat:", "Planul de verde poate fi programat direct în automatele de dirijare din teren.", font_size=12, space_after=0)
    
    # Bottom Transparency Card
    top_t = Inches(6.05)
    h_t = Inches(0.90)
    w_full = Inches(11.73)
    add_card(slide, Inches(0.8), top_t, w_full, h_t, bg_color=C_CARD_HI, border_color=C_BORDER)
    tf_t = add_textbox(slide, Inches(0.8), top_t, w_full, h_t, margin=0.15)
    add_p(tf_t, "TRANSPARENȚĂ COMPLETĂ: CE E REAL ȘI CE E SIMULAT", font_size=11, bold=True, color=C_CYAN, space_after=3)
    add_p(tf_t, "100% REAL: Rețeaua rutieră OSMnx Iași, formulele de optimizare Webster, calculele de capacitate și aplicațiile Web/Mobile.\nSIMULAT CU FIDELITATE: Volumele microscopice de trafic generate în SUMO pe baza distribuțiilor orare oficiale din PMUD Iași.", 
          font_size=11, color=C_SLATE_LIGHT, space_after=0)

# ==========================================
# SLIDE 8: PITCH 5 · DEMO LIVE
# ==========================================
def build_slide_8(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · DEMONSTRAȚIE LIVE", "Demo Live: Optimizarea Nodului Podu Roșu", 
               "Trei pași simpli de la identificarea blocajului la scenariul optimizat testat în timp real pe ecran și pe mobil.", "05 / 10 · DEMO")
    
    # Top Viewport Card
    top_v = Inches(1.85)
    h_v = Inches(2.45)
    w_full = Inches(11.73)
    
    add_card(slide, Inches(0.8), top_v, w_full, h_v, bg_color=C_CARD_HI, border_color=C_BORDER_HI, border_width=1.5)
    tf_v = add_textbox(slide, Inches(0.8), top_v, w_full, h_v, margin=0.22)
    add_p(tf_v, "ARTERIA CONTROL CENTER · IAȘI DIGITAL TWIN [SIMULARE ACTIVĂ]", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tf_v, "Vizualizare Interactivă MapLibre GL · Coridor Podu Roșu – Bd. Socola", font_size=17, bold=True, color=C_WHITE, space_after=12)
    
    # 3 Metric boxes inside viewport
    add_bullet_run(tf_v, "• ÎNTÂRZIERE MEDIE PE CORIDOR:", "85 secunde  ➔  56 secunde (-34% reducere timp)", font_size=13, space_after=6, text_color=C_EMERALD)
    add_bullet_run(tf_v, "• LUNGIME COADĂ BD. SOCOLA:", "340 metri  ➔  190 metri (-44% reducere coadă)", font_size=13, space_after=6, text_color=C_EMERALD)
    add_bullet_run(tf_v, "• VITEZĂ MEDIE ÎN ORA DE VÂRF:", "14 km/h  ➔  21 km/h (+50% creștere viteză)", font_size=13, space_after=0, text_color=C_EMERALD)
    
    # Bottom 3 Step Cards
    top_s = Inches(4.45)
    h_s = Inches(2.50)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Step 1
    add_card(slide, Inches(0.8), top_s, w_c, h_s, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_s, w_c, h_s, margin=0.20)
    add_p(tf1, "PASUL 1: SELECTARE SCENARIU", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_bullet_run(tf1, "• Nod Critic:", "Podu Roșu cu 6 intrări rutiere.", font_size=12, space_after=6)
    add_bullet_run(tf1, "• Profil:", "„Vârf Dimineață (08:00–09:00)”.", font_size=12, space_after=6)
    add_bullet_run(tf1, "• Diagnostic:", "Flux masiv 2.180 veh/h spre centru; blocaj pe banda de dreapta Socola.", font_size=12, space_after=0)
    
    # Step 2
    add_card(slide, Inches(0.8) + w_c + gap, top_s, w_c, h_s, bg_color=C_CARD, border_color=C_BORDER)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_s, w_c, h_s, margin=0.20)
    add_p(tf2, "PASUL 2: RULARE SIMULARE SUMO", font_size=11, bold=True, color=C_AMBER, space_after=6)
    add_bullet_run(tf2, "• Pornire TraCI:", "Un click declanșează micro-simularea.", font_size=12, space_after=6)
    add_bullet_run(tf2, "• Animație Vehicule:", "Traiectorii fizice individuale pe benzi.", font_size=12, space_after=6)
    add_bullet_run(tf2, "• Validare:", "Identifică exact momentul în care coada blochează accesul ambulanțelor.", font_size=12, space_after=0)
    
    # Step 3
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_s, w_c, h_s, bg_color=C_CARD_ELEV, border_color=C_BORDER_EM, border_width=1.2)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_s, w_c, h_s, margin=0.20)
    add_p(tf3, "PASUL 3: APLICARE MODEL WEBSTER", font_size=11, bold=True, color=C_EMERALD, space_after=6)
    add_bullet_run(tf3, "• Optimizare:", "Recalculare ciclu optim la 110 secunde.", font_size=12, space_after=6)
    add_bullet_run(tf3, "• Redistribuire:", "+14s verde pe axa aglomerată.", font_size=12, space_after=6)
    add_bullet_run(tf3, "• Rezultat:", "Comparație Before/After instant pe ecran și în aplicația mobilă Expo.", font_size=12, space_after=0)

# ==========================================
# SLIDE 9: PITCH 6 · CUM E CONSTRUIT
# ==========================================
def build_slide_9(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · ARHITECTURĂ & DECIZII TEHNICE", "Arhitectura Tehnică și Deciziile de Sistem", 
               "Arhitectură decuplată: calcul analitic instant (<1s) combinat cu validare microscopică fizică în fundal.", "06 / 10 · ARHITECTURĂ")
    
    # Top 4 Pipeline Cards
    top_p = Inches(1.85)
    h_p = Inches(2.15)
    w_p = Inches(2.68)
    gap_p = Inches(0.33)
    
    stages = [
        ("1. INGESTIE GEOMETRIE", C_CYAN, "OSMnx & Overpass", 
         "Extrage automat topologia străzilor, benzi, priorități și sensuri giratorii din OpenStreetMap pentru Iași."),
        ("2. COMPUTE ENGINES", C_EMERALD, "FastAPI + SUMO + Webster", 
         "Model analitic Webster (<1s) pentru UI + Eclipse SUMO 1.28.0 (TraCI) pentru micro-simulare fizică detaliată."),
        ("3. STATE & DATA VALIDATION", C_CYAN, "Pydantic v2 + Caching", 
         "Scheme stricte de validare a datelor, serializare GeoJSON standard și gestionare a istoricului scenariilor."),
        ("4. CLIENȚI MULTIPLATFORMĂ", C_AMBER, "Next.js 14 & Expo Android", 
         "Rendare GPU cu MapLibre GL pe Web și aplicație nativă React Native pentru inspecția pe teren a nodurilor.")
    ]
    
    for i, (title_s, col_s, tech_s, desc_s) in enumerate(stages):
        x = Inches(0.8) + i * (w_p + gap_p)
        add_card(slide, x, top_p, w_p, h_p, bg_color=C_CARD, border_color=C_BORDER)
        tf = add_textbox(slide, x, top_p, w_p, h_p, margin=0.16)
        add_p(tf, title_s, font_size=10, bold=True, color=col_s, space_after=4)
        add_p(tf, tech_s, font_size=13, bold=True, color=C_WHITE, space_after=6)
        add_p(tf, desc_s, font_size=11, color=C_SLATE_LIGHT, space_after=0)
        
    # Bottom 2 Decision Cards
    top_d = Inches(4.18)
    h_d = Inches(2.77)
    w_d = Inches(5.68)
    gap_d = Inches(0.37)
    
    # Decision 1: Hybrid Engine
    add_card(slide, Inches(0.8), top_d, w_d, h_d, bg_color=C_CARD_ELEV, border_color=C_BORDER)
    tf_d1 = add_textbox(slide, Inches(0.8), top_d, w_d, h_d, margin=0.22)
    add_p(tf_d1, "DECIZIA CHEIE: ARHITECTURĂ HIBRIDĂ WEBSTER + SUMO", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf_d1, "De ce două motoare de calcul?", font_size=15, bold=True, color=C_WHITE, space_after=8)
    add_bullet_run(tf_d1, "• Problema SUMO:", "O simulare microscopică completă TraCI pe coridor durează 30-60 de secunde pe server. Într-o ședință de lucru la primărie, utilizatorul nu poate aștepta un minut la fiecare click.", font_size=12, space_after=8)
    add_bullet_run(tf_d1, "• Soluția Noastră:", "Modelul matematic analitic Webster calculează soluția optimă în <1 secundă direct în interfață, în timp ce SUMO rulează asincron în fundal pentru confirmarea microscopică fizică.", font_size=12, space_after=0, text_color=C_EMERALD)
    
    # Decision 2: AI vs Human
    add_card(slide, Inches(0.8) + w_d + gap_d, top_d, w_d, h_d, bg_color=C_CARD_HI, border_color=C_BORDER)
    tf_d2 = add_textbox(slide, Inches(0.8) + w_d + gap_d, top_d, w_d, h_d, margin=0.22)
    add_p(tf_d2, "TRANSPARENȚĂ: CE A FĂCUT AI vs. CE AM DEȚINUT NOI", font_size=11, bold=True, color=C_AMBER, space_after=6)
    add_p(tf_d2, "Unde a fost folosit AI-ul și unde e meritul echipei?", font_size=15, bold=True, color=C_WHITE, space_after=8)
    add_bullet_run(tf_d2, "• Accelerat de AI (Copilot/Gemini):", "Generarea de boilerplate Python/TypeScript, scheme Pydantic repetitive, conversie de stiluri CSS și mock-uri GeoJSON inițiale.", font_size=12, space_after=8)
    add_bullet_run(tf_d2, "• Deciziile Noastre (Iustin & Alex):", "Arhitectura hibridă de simulare, calibrarea parametrilor de trafic pe Podu Roșu, logica de prioritizare a transportului public și design-ul de produs.", font_size=12, space_after=0, text_color=C_WHITE)

# ==========================================
# SLIDE 10: PITCH 7 · BUSINESS / CINE PLĂTEȘTE
# ==========================================
def build_slide_10(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · MODEL ECONOMIC & SUSTENABILITATE", "Model Economic: B2G SaaS pentru Orașe Medii", 
               "Preț calibrat sub pragul de achiziție directă pentru a asigura o vânzare rapidă în 14 zile, fără licitații de 18 luni.", "07 / 10 · MODEL ECONOMIC")
    
    top_c = Inches(1.85)
    h_c = Inches(5.10)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Col 1: Who pays
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.22)
    add_p(tf1, "CINE PLĂTEȘTE (CLIENTUL B2G)", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf1, "Administrația Locală", font_size=18, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf1, "• Primăria Municipiului Iași:", "Direcția Tehnică, Serviciul Siguranța Circulației și Managementul Traficului.", font_size=12, space_after=10)
    add_bullet_run(tf1, "• CTP Iași (Transport Public):", "Interes direct de a reduce timpii de parcurs pe liniile cheie de autobuz și tramvai.", font_size=12, space_after=10)
    add_bullet_run(tf1, "• Justificare Bugetară:", "Fonduri dedicate de mobilitate urbană inteligentă (PNRR / POR), unde se cer justificări matematice clare.", font_size=12, space_after=0)
    
    # Col 2: Pricing Structure
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER_EM, border_width=1.5)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.22)
    add_p(tf2, "STRUCTURA DE PREȚ (B2G SAAS)", font_size=11, bold=True, color=C_EMERALD, space_after=6)
    add_p(tf2, "1.500€ / lună", font_size=34, bold=True, color=C_EMERALD, space_after=0)
    add_p(tf2, "18.000€ pe an per municipiu", font_size=12, bold=True, color=C_WHITE, space_after=14)
    add_bullet_run(tf2, "• Sub Pragul de Achiziție Directă:", "În România, pragul de achiziție directă este de 270.120 lei (~54.000€). Contractul se semnează direct, fără licitații de 18 luni.", font_size=12, space_after=10, text_color=C_ICE)
    add_bullet_run(tf2, "• Pachetul Include:", "Simulare cloud nelimitată, actualizări săptămânale de geometrie OSMnx, suport dedicat și export rapoarte oficiale PMUD.", font_size=12, space_after=0)
    
    # Col 3: Unit Economics
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.22)
    add_p(tf3, "UNIT ECONOMICS & MARJĂ", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf3, "> 95%", font_size=34, bold=True, color=C_WHITE, space_after=0)
    add_p(tf3, "marjă brută per contract anual", font_size=11, color=C_SLATE_MUTED, space_after=14)
    add_bullet_run(tf3, "• Cost Cloud Redus:", "~45€ / lună per oraș (instanțe CPU la cerere pentru FastAPI + SUMO TraCI) = ~540€ / an.", font_size=12, space_after=10)
    add_bullet_run(tf3, "• ROI Masiv pentru Oraș:", "Un singur studiu clasic de trafic costă 50.000€. Cu ARTERIA, primăria economisește 32.000€ din start și are simulare continuă.", font_size=12, space_after=0, text_color=C_EMERALD)

# ==========================================
# SLIDE 11: PITCH 8 · PIAȚA
# ==========================================
def build_slide_11(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · DIMENSIUNEA PIEȚEI", "Piața: Calculată de Jos în Sus (Bottom-Up)", 
               "De la validarea pe coridorul Podu Roșu din Iași la cele 40 de municipii reședință de județ și regiunea CEE.", "08 / 10 · PIAȚA")
    
    top_c = Inches(1.85)
    h_m = Inches(2.25)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Metric 1: SOM
    add_card(slide, Inches(0.8), top_c, w_c, h_m, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_m, margin=0.20)
    add_p(tf1, "SOM — PIAȚA OBȚINABILĂ IMEDIAT", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tf1, "18.000€ ARR", font_size=28, bold=True, color=C_WHITE, space_after=2)
    add_p(tf1, "Municipiul Iași (Primărie + CTP)", font_size=12, bold=True, color=C_ICE, space_after=4)
    add_p(tf1, "Validare operațională pe coridorul pilot Podu Roșu – Socola ca primă referință B2G.", font_size=11, color=C_SLATE_LIGHT, space_after=0)
    
    # Metric 2: SAM
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_m, bg_color=C_CARD_ELEV, border_color=C_BORDER_EM, border_width=1.2)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_m, margin=0.20)
    add_p(tf2, "SAM — PIAȚA ADRESABILĂ ROMÂNIA", font_size=11, bold=True, color=C_EMERALD, space_after=4)
    add_p(tf2, "720.000€ ARR", font_size=28, bold=True, color=C_EMERALD, space_after=2)
    add_p(tf2, "40 de Municipii Reședință de Județ", font_size=12, bold=True, color=C_WHITE, space_after=4)
    add_p(tf2, "40 orașe × 18.000€/an. Toate au obligația legală de a deține și actualiza planurile de mobilitate PMUD.", font_size=11, color=C_SLATE_LIGHT, space_after=0)
    
    # Metric 3: TAM
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_m, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_m, margin=0.20)
    add_p(tf3, "TAM — EUROPA CENTRALĂ ȘI DE EST", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tf3, "14.4M€ ARR", font_size=28, bold=True, color=C_WHITE, space_after=2)
    add_p(tf3, "800+ Municipii Medii în CEE", font_size=12, bold=True, color=C_ICE, space_after=4)
    add_p(tf3, "Orașe între 50k și 500k locuitori din Polonia, Cehia, Ungaria, Bulgaria și România cu probleme similare.", font_size=11, color=C_SLATE_LIGHT, space_after=0)
    
    # Bottom 3 Expansion Steps
    top_s = Inches(4.30)
    h_s = Inches(2.65)
    
    add_card(slide, Inches(0.8), top_s, w_c, h_s, bg_color=C_CARD, border_color=C_BORDER)
    tfs1 = add_textbox(slide, Inches(0.8), top_s, w_c, h_s, margin=0.20)
    add_p(tfs1, "FAZA 1 (Q1–Q2 2026): PILOTUL IAȘI", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_bullet_run(tfs1, "• Acord de Pilot:", "Semnare parteneriat formal cu Primăria Iași.", font_size=12, space_after=6)
    add_bullet_run(tfs1, "• Testare Linia 28:", "Validare scădere întârzieri pe transportul public.", font_size=12, space_after=6)
    add_bullet_run(tfs1, "• Studiu de Caz:", "Publicarea primelor date verificate de teren.", font_size=12, space_after=0)
    
    add_card(slide, Inches(0.8) + w_c + gap, top_s, w_c, h_s, bg_color=C_CARD, border_color=C_BORDER)
    tfs2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_s, w_c, h_s, margin=0.20)
    add_p(tfs2, "FAZA 2 (Q3–Q4 2026): POLII REGIONALI", font_size=11, bold=True, color=C_EMERALD, space_after=6)
    add_bullet_run(tfs2, "• Extindere Națională:", "Cluj-Napoca, Timișoara, Brașov, Oradea.", font_size=12, space_after=6)
    add_bullet_run(tfs2, "• Target Financiar:", "5 contracte B2G active (90.000€ ARR).", font_size=12, space_after=6)
    add_bullet_run(tfs2, "• Vânzare Directă:", "Contractare directă sub pragul legal de achiziție.", font_size=12, space_after=0)
    
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_s, w_c, h_s, bg_color=C_CARD, border_color=C_BORDER)
    tfs3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_s, w_c, h_s, margin=0.20)
    add_p(tfs3, "FAZA 3 (2027): SCALARE REGIONALĂ CEE", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_bullet_run(tfs3, "• Import Automat:", "Localizare instantă a oricărui oraș din CEE via OSMnx.", font_size=12, space_after=6)
    add_bullet_run(tfs3, "• Parteneriate ITS:", "Integrare cu marii distribuitori de echipamente de semaforizare.", font_size=12, space_after=6)
    add_bullet_run(tfs3, "• Obiectiv:", "25 de orașe active în rețea.", font_size=12, space_after=0)

# ==========================================
# SLIDE 12: PITCH 9 · ECHIPA
# ==========================================
def build_slide_12(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · ECHIPA FONDATOARE", "Echipa ARION: Cine a Deținut Ce & De Ce Noi", 
               "O echipă tehnică locală, complementară, dedicată rezolvării blocajului în care stăm în fiecare zi.", "09 / 10 · ECHIPA")
    
    # 4 Team Member Cards across full width
    top_t = Inches(1.85)
    h_t = Inches(3.30)
    w_t = Inches(2.68)
    gap_t = Inches(0.33)
    
    # Card 1: Iustin Fagateanu
    add_card(slide, Inches(0.8), top_t, w_t, h_t, bg_color=C_CARD_ELEV, border_color=C_BORDER_HI, border_width=1.5)
    tf1 = add_textbox(slide, Inches(0.8), top_t, w_t, h_t, margin=0.18)
    add_p(tf1, "IUSTIN FAGATEANU", font_size=14, bold=True, color=C_WHITE, space_after=2)
    add_p(tf1, "Product & Frontend Lead", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_bullet_run(tf1, "• A conceput arhitectura Web:", "Next.js 14 cu hartă interactivă MapLibre GL 3D.", font_size=11, space_after=5)
    add_bullet_run(tf1, "• A dezvoltat clientul mobil:", "Aplicație Expo Android pentru echipele de teren.", font_size=11, space_after=5)
    add_bullet_run(tf1, "• A proiectat sistemul UX:", "Interfață intuitivă pentru compararea scenariilor A/B de trafic.", font_size=11, space_after=0)
    
    # Card 2: Alex Pintilie
    add_card(slide, Inches(0.8) + w_t + gap_t, top_t, w_t, h_t, bg_color=C_CARD_ELEV, border_color=C_BORDER_HI, border_width=1.5)
    tf2 = add_textbox(slide, Inches(0.8) + w_t + gap_t, top_t, w_t, h_t, margin=0.18)
    add_p(tf2, "ALEX PINTILIE", font_size=14, bold=True, color=C_WHITE, space_after=2)
    add_p(tf2, "Backend & Simulation Lead", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_bullet_run(tf2, "• A construit serverul FastAPI:", "Arhitectură asincronă de procesare a fluxurilor de simulare.", font_size=11, space_after=5)
    add_bullet_run(tf2, "• A integrat motorul SUMO:", "Conexiune TraCI pentru rularea microscopică fizică.", font_size=11, space_after=5)
    add_bullet_run(tf2, "• A implementat modelul Webster:", "Algoritm de calcul analitic al timpilor optimi de verde.", font_size=11, space_after=0)
    
    # Card 3: Modelare Urbana
    add_card(slide, Inches(0.8) + (w_t + gap_t)*2, top_t, w_t, h_t, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_t + gap_t)*2, top_t, w_t, h_t, margin=0.18)
    add_p(tf3, "ECHIPA ARION", font_size=14, bold=True, color=C_WHITE, space_after=2)
    add_p(tf3, "Modelare Urbană & Trafic", font_size=11, bold=True, color=C_EMERALD, space_after=6)
    add_bullet_run(tf3, "• Calibrare date reale:", "Integrare volume PMUD Iași și măsurători senzori la Podu Roșu.", font_size=11, space_after=5)
    add_bullet_run(tf3, "• Topologie rețea:", "Configurarea nodului Podu Roșu cu 6 intrări rutiere în OSMnx.", font_size=11, space_after=5)
    add_bullet_run(tf3, "• Evaluare Before / After:", "Măsurarea reducerii cozilor și timpilor de întârziere.", font_size=11, space_after=0)
    
    # Card 4: Advisory
    add_card(slide, Inches(0.8) + (w_t + gap_t)*3, top_t, w_t, h_t, bg_color=C_CARD, border_color=C_BORDER)
    tf4 = add_textbox(slide, Inches(0.8) + (w_t + gap_t)*3, top_t, w_t, h_t, margin=0.18)
    add_p(tf4, "ADVISORY MOBILITATE", font_size=14, bold=True, color=C_WHITE, space_after=2)
    add_p(tf4, "Domeniu & Administrație", font_size=11, bold=True, color=C_AMBER, space_after=6)
    add_bullet_run(tf4, "• Consultare transport public:", "Validarea cerințelor operaționale ale CTP Iași pentru tramvaie/autobuze.", font_size=11, space_after=5)
    add_bullet_run(tf4, "• Cerințe tehnice primărie:", "Aliniere cu atribuțiile direcțiilor de administrare a circulației.", font_size=11, space_after=5)
    add_bullet_run(tf4, "• Conformitate B2G:", "Verificarea încadrării în cerințele legale de achiziție directă.", font_size=11, space_after=0)
    
    # Bottom 2 Commitment Cards
    top_c = Inches(5.30)
    h_c = Inches(1.65)
    w_half = Inches(5.68)
    gap_h = Inches(0.37)
    
    add_card(slide, Inches(0.8), top_c, w_half, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tfc1 = add_textbox(slide, Inches(0.8), top_c, w_half, h_c, margin=0.18)
    add_p(tfc1, "DE CE NOI PENTRU ACEASTĂ PROBLEMĂ", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tfc1, "Locuim în Iași, ne confruntăm zilnic cu blocajul de la Podu Roșu și am demonstrat în 48 de ore că putem uni tehnologii moderne (FastAPI, SUMO, Next.js, GIS) pe care firmele clasice de consultanță nu le stăpânesc.", 
          font_size=11, color=C_SLATE_LIGHT, space_after=0)
    
    add_card(slide, Inches(0.8) + w_half + gap_h, top_c, w_half, h_c, bg_color=C_CARD_HI, border_color=C_BORDER_HI)
    tfc2 = add_textbox(slide, Inches(0.8) + w_half + gap_h, top_c, w_half, h_c, margin=0.18)
    add_p(tfc2, "CINE CONTINUĂ DUPĂ HACKATHON", font_size=11, bold=True, color=C_EMERALD, space_after=4)
    add_p(tfc2, "Ambii fondatori tehnici (Iustin Fagateanu & Alex Pintilie) sunt 100% dedicați pentru continuarea proiectului și livrarea primului pilot operațional în parteneriat cu Primăria Municipiului Iași și CTP.", 
          font_size=11, bold=True, color=C_WHITE, space_after=0)

# ==========================================
# SLIDE 13: PITCH 10 · CE URMEAZĂ & CEREREA NOASTRĂ
# ==========================================
def build_slide_13(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · PASUL URMĂTOR & CEREREA NOASTRĂ", "Ce Urmează: Pilotul Operațional și Cererea Noastră", 
               "Planul clar de după scenă: ce livrăm în 30 de zile, ce măsurăm în teren și de ce avem nevoie din partea juriului.", "10 / 10 · CE URMEAZĂ")
    
    top_c = Inches(1.85)
    h_c = Inches(3.60)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Col 1
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.20)
    add_p(tf1, "1. AM TĂIAT ȘI ADĂUGĂM ÎN Q2", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf1, "Dezvoltare Tehnică", font_size=16, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf1, "• Analiză Video Automată:", "Detecție automată a fluxurilor de vehicule din camerele de supraveghere existente în intersecții.", font_size=12, space_after=8)
    add_bullet_run(tf1, "• Export Direct GIS/CAD:", "Compatibilitate nativă cu formatele utilizate în birourile de urbanism.", font_size=12, space_after=8)
    add_bullet_run(tf1, "• Prioritizare Dinamică CTP:", "Extensie automată de verde la apropierea tramvaielor întârziate.", font_size=12, space_after=0)
    
    # Col 2
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER_EM, border_width=1.2)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.20)
    add_p(tf2, "2. CE MĂSURĂM PRIMUL ÎN TEREN", font_size=11, bold=True, color=C_EMERALD, space_after=6)
    add_p(tf2, "Validare Metrică în 30 Zile", font_size=16, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf2, "• Linia 28 CTP Iași:", "Măsurăm timpii de parcurs pe tronsonul Socola – Podu Roșu – Centru.", font_size=12, space_after=8)
    add_bullet_run(tf2, "• Validare Matematică:", "Verificăm dacă reducerea simulată de 34% generează cel puțin 20% reducere pe stradă.", font_size=12, space_after=8, text_color=C_EMERALD)
    add_bullet_run(tf2, "• Raport Public B2G:", "Predarea către primărie a primului dosar tehnic bazat pe date verificate.", font_size=12, space_after=0)
    
    # Col 3
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD_HI, border_color=C_BORDER_HI, border_width=1.2)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.20)
    add_p(tf3, "3. CEREREA NOASTRĂ CONCRETĂ", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf3, "Ce Solicităm Juriului", font_size=16, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf3, "• Acord de Pilot B2G:", "Deschiderea dialogului instituțional cu Primăria Iași și Direcția Tehnică.", font_size=12, space_after=8)
    add_bullet_run(tf3, "• Diagrame Semaforizare:", "Acces la planurile reale de semaforizare din intersecția Podu Roșu.", font_size=12, space_after=8)
    add_bullet_run(tf3, "• Mentorat Achiziții Publice:", "Îndrumare juridică pentru structurarea contractului de achiziție directă.", font_size=12, space_after=0)
    
    # Bottom Closing Statement
    top_close = Inches(5.60)
    h_close = Inches(1.35)
    w_full = Inches(11.73)
    
    add_card(slide, Inches(0.8), top_close, w_full, h_close, bg_color=C_CARD_ELEV, border_color=C_BORDER_HI)
    tf_c = add_textbox(slide, Inches(0.8), top_close, w_full, h_close, margin=0.18)
    add_p(tf_c, "„Nu putem lărgi străzile Iașului cu încă o bandă, dar putem transforma fiecare secundă de verde într-o decizie calculată matematic.”", 
          font_size=15, bold=True, italic=True, color=C_WHITE, align=PP_ALIGN.CENTER, space_after=8)
    add_p(tf_c, "contact@arteria.ro   ·   arteria.ai-iasi.ro   ·   Echipa ARION (Iustin Fagateanu & Alex Pintilie)", 
          font_size=12, bold=True, color=C_CYAN, align=PP_ALIGN.CENTER, space_after=0)

# ==========================================
# SLIDE 14: ANEXE
# ==========================================
def build_slide_14(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · ANEXE TEHNICE & DOCUMENTAȚIE", "Slide-uri de Suport pentru Sesiunea de Q&A", 
               "Detalii concrete care răspund dinainte la întrebările tehnice, metodologice și de business ale juriului.", "ANEXE · SUPORT")
    
    top_c = Inches(1.85)
    h_c = Inches(5.10)
    w_c = Inches(2.68)
    gap = Inches(0.33)
    
    # Anexa A
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.18)
    add_p(tf1, "ANEXA A", font_size=11, bold=True, color=C_CORAL, space_after=4)
    add_p(tf1, "Ce Nu Merge Încă", font_size=15, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf1, "• Rulare SUMO Asincronă:", "Simularea completă a întregului oraș durează 40-60s pe server; de aceea pe Web folosim modelul instant Webster.", font_size=11, space_after=8)
    add_bullet_run(tf1, "• Date Pietonale OSM:", "OpenStreetMap nu conține lățimea exactă a tuturor trotuarelor pe străzile secundare din Iași.", font_size=11, space_after=8)
    add_bullet_run(tf1, "• Roadmap:", "Cluster distribuit de simulare și cartografiere crowdsourced Mapillary în faza a doua.", font_size=11, space_after=0)
    
    # Anexa B
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.18)
    add_p(tf2, "ANEXA B", font_size=11, bold=True, color=C_CYAN, space_after=4)
    add_p(tf2, "Detalii Tehnice Stack", font_size=15, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf2, "• Frontend Web:", "Next.js 14 + MapLibre GL 3D + Tailwind CSS + shadcn/ui.", font_size=11, space_after=8)
    add_bullet_run(tf2, "• Mobile Client:", "Expo / React Native pentru Android.", font_size=11, space_after=8)
    add_bullet_run(tf2, "• Backend API:", "FastAPI + Python 3.13 + Pydantic v2 strict.", font_size=11, space_after=8)
    add_bullet_run(tf2, "• Micro-Simulare:", "Eclipse SUMO 1.28.0 TraCI + OSMnx 1.9.", font_size=11, space_after=0)
    
    # Anexa C
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.18)
    add_p(tf3, "ANEXA C", font_size=11, bold=True, color=C_AMBER, space_after=4)
    add_p(tf3, "Dovezile Complete", font_size=15, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf3, "• Rapoarte Oficiale:", "PMUD Iași (2023) · TomTom Traffic Index Report (2024) · Planul Integrat de Calitate a Aerului Iași (2024).", font_size=11, space_after=8)
    add_bullet_run(tf3, "• Comparație Costuri:", "Consultanță clasică (50k€, 18 luni) vs. SCADA rigid (1M€+, fără model GIS) vs. ARTERIA (18k€/an SaaS, 5 secunde).", font_size=11, space_after=0)
    
    # Anexa D
    add_card(slide, Inches(0.8) + (w_c + gap)*3, top_c, w_c, h_c, bg_color=C_CARD_HI, border_color=C_BORDER_HI)
    tf4 = add_textbox(slide, Inches(0.8) + (w_c + gap)*3, top_c, w_c, h_c, margin=0.18)
    add_p(tf4, "ANEXA D", font_size=11, bold=True, color=C_EMERALD, space_after=4)
    add_p(tf4, "Calculele Financiare", font_size=15, bold=True, color=C_WHITE, space_after=10)
    add_bullet_run(tf4, "• Calcul SAM:", "40 orașe reședință × 1.500€/lună × 12 luni = 720.000€ ARR în România.", font_size=11, space_after=8)
    add_bullet_run(tf4, "• Achiziție Directă:", "18.000€/an este mult sub pragul legal din România (270.120 RON = ~54.000€).", font_size=11, space_after=8)
    add_bullet_run(tf4, "• Cost Server:", "~540€/an per oraș -> Marjă brută > 95%.", font_size=11, space_after=0, text_color=C_EMERALD)

# ==========================================
# SLIDE 15: Q&A
# ==========================================
def build_slide_15(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · SESIUNEA DE Q&A", "Șapte Întrebări Cheie & Răspunsurile Noastre Clare", 
               "Suntem pregătiți să apărăm fiecare decizie tehnică, arhitecturală și economică din fața juriului.", "Q&A · PREGĂTIRE")
    
    top_c = Inches(1.85)
    h_row = Inches(0.66)
    gap_r = Inches(0.08)
    w_full = Inches(11.73)
    
    qa_list = [
        ("„Cine e utilizatorul real?”", "Inginerul din Direcția Tehnică a Primăriei Iași și dispeceratul CTP Iași (operatorul de transport).", "Slide 02 & 07", C_CYAN),
        ("„Există deja alternative?”", "Da: studii statice lente (50k€, 18 luni) și sisteme SCADA de 1M€. Noi oferim simulare instant la 18k€/an.", "Slide 03", C_WHITE),
        ("„Cât din cod este scris de AI?”", "Boilerplate-ul și mock-urile; arhitectura hibridă, modelul Webster și integrarea SUMO sunt 100% ale noastre.", "Slide 06", C_AMBER),
        ("„Ce nu merge încă în produs?”", "Simularea pe tot municipiul durează 40-60s în SUMO; datele pietonale OSM pe străzi secundare sunt incomplete.", "Anexa A", C_CORAL),
        ("„Ce măsurați concret în pilot?”", "Timpul de parcurs al autobuzului 28 CTP pe tronsonul Socola – Podu Roșu în 30 de zile de pilot.", "Slide 10", C_EMERALD),
        ("„Cine plătește și din ce fonduri?”", "Primăria Iași prin achiziție directă (18.000€/an, sub pragul legal de 54.000€), din bugete de mobilitate inteligentă.", "Slide 07", C_WHITE),
        ("„De ce voi și ce faceți de luni?”", "Iustin Fagateanu & Alex Pintilie: echipă tehnică complementară, continuăm full-time pentru pilotul Podu Roșu.", "Slide 09", C_CYAN)
    ]
    
    for i, (q, a, s_ref, c_q) in enumerate(qa_list):
        y = top_c + i * (h_row + gap_r)
        add_card(slide, Inches(0.8), y, w_full, h_row, bg_color=C_CARD, border_color=C_BORDER)
        tf = add_textbox(slide, Inches(0.8), y, w_full, h_row, margin=0.12)
        p = tf.paragraphs[0]
        p.space_after = Pt(0)
        
        # Q
        r1 = p.add_run()
        r1.text = q + "  ➔  "
        r1.font.name = 'Calibri'
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = c_q
        
        # A
        r2 = p.add_run()
        r2.text = a + "   "
        r2.font.name = 'Calibri'
        r2.font.size = Pt(12)
        r2.font.color.rgb = C_SLATE_LIGHT
        
        # Ref
        r3 = p.add_run()
        r3.text = f"[{s_ref}]"
        r3.font.name = 'Calibri'
        r3.font.size = Pt(11)
        r3.font.bold = True
        r3.font.color.rgb = C_SLATE_DARK

# ==========================================
# SLIDE 16: ULTIMUL LUCRU / CHECKLIST
# ==========================================
def build_slide_16(slide):
    clear_slide(slide)
    add_bg(slide)
    add_header(slide, "ARTERIA · LISTA DE VERIFICARE", "Ultimul Lucru: Pregătirea Înainte de Scenă", 
               "Checklist-ul esențial înainte de pitch: conținut validat, reguli de scenă și ce nu facem niciodată.", "CHECKLIST FINAL")
    
    top_c = Inches(1.85)
    h_c = Inches(5.10)
    w_c = Inches(3.68)
    gap = Inches(0.34)
    
    # Col 1: Conținut
    add_card(slide, Inches(0.8), top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf1 = add_textbox(slide, Inches(0.8), top_c, w_c, h_c, margin=0.22)
    add_p(tf1, "CONȚINUT VALIDAT", font_size=11, bold=True, color=C_EMERALD, space_after=6)
    add_p(tf1, "Mesaj Clar & Cifre Reale", font_size=16, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf1, "✓ Problema într-o frază:", "Om concret (inginerul de trafic), loc (Podu Roșu) și cost (42 min/zi).", font_size=12, space_after=10)
    add_bullet_run(tf1, "✓ Surse oficiale pe slide:", "PMUD Iași (2023), senzori trafic și TomTom Traffic Index 2024.", font_size=12, space_after=10)
    add_bullet_run(tf1, "✓ Real vs. Simulat:", "Declarat deschis și transparent, fără exagerări tehnice.", font_size=12, space_after=10)
    add_bullet_run(tf1, "✓ Cererea concretă:", "Acord de pilot pe Podu Roșu din partea Primăriei Iași și acces la senzori.", font_size=12, space_after=0)
    
    # Col 2: Repetiție
    add_card(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, bg_color=C_CARD_ELEV, border_color=C_BORDER_HI)
    tf2 = add_textbox(slide, Inches(0.8) + w_c + gap, top_c, w_c, h_c, margin=0.22)
    add_p(tf2, "REGULI PENTRU SCENĂ", font_size=11, bold=True, color=C_CYAN, space_after=6)
    add_p(tf2, "Execuție Fără Greșeală", font_size=16, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf2, "✓ Demo testat offline:", "Înregistrare video de rezervă pregătită în caz de cădere internet.", font_size=12, space_after=10)
    add_bullet_run(tf2, "✓ Cronometru strict:", "Pitch repetat în 2:45 min (rezervă de 15 secunde pentru siguranță).", font_size=12, space_after=10)
    add_bullet_run(tf2, "✓ Roluri clare:", "Iustin Fagateanu prezintă produsul & viziunea; Alex Pintilie răspunde la arhitectură și SUMO.", font_size=12, space_after=10)
    add_bullet_run(tf2, "✓ Păstrare calm:", "La întrebări grele spunem „Iată cum validăm în pilot”, fără defensivitate.", font_size=12, space_after=0)
    
    # Col 3: Ce nu faci
    add_card(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, bg_color=C_CARD, border_color=C_BORDER)
    tf3 = add_textbox(slide, Inches(0.8) + (w_c + gap)*2, top_c, w_c, h_c, margin=0.22)
    add_p(tf3, "CE NU FACEM NICIODATĂ", font_size=11, bold=True, color=C_CORAL, space_after=6)
    add_p(tf3, "Greșelile Care Costă Puncte", font_size=16, bold=True, color=C_WHITE, space_after=12)
    add_bullet_run(tf3, "✗ Nu începe cu echipa:", "Juriul vrea să știe întâi ce problemă rezolvi și pe cine ajuți.", font_size=12, space_after=10, label_color=C_CORAL)
    add_bullet_run(tf3, "✗ Nu citi slide-urile:", "Textul e scheletul; tu aduci energia, contextul și povestea.", font_size=12, space_after=10, label_color=C_CORAL)
    add_bullet_run(tf3, "✗ Nu spune «totul merge»:", "Arată maturitate numind deschis ce mai e de optimizat în simulator.", font_size=12, space_after=10, label_color=C_CORAL)
    add_bullet_run(tf3, "✗ Nu încheia cu «Mulțumim»:", "Ultimul cuvânt este cererea de pilot și viziunea matematică.", font_size=12, space_after=0, label_color=C_CORAL)


# ==========================================
# MAIN EXECUTION
# ==========================================
def main():
    prs_path = 'docs/ppt/Template-Pitch-Hackathon.pptx'
    prs = pptx.Presentation(prs_path)
    print(f"Loaded presentation: {prs_path} ({len(prs.slides)} slides)")
    
    # Slide 1: Cover
    print("Building Slide 1...")
    build_slide_1(prs.slides[0])
    
    # Slide 4 to 13: Core Pitch Slides
    print("Building Slide 4 (Pitch 1)...")
    build_slide_4(prs.slides[3])
    
    print("Building Slide 5 (Pitch 2)...")
    build_slide_5(prs.slides[4])
    
    print("Building Slide 6 (Pitch 3)...")
    build_slide_6(prs.slides[5])
    
    print("Building Slide 7 (Pitch 4)...")
    build_slide_7(prs.slides[6])
    
    print("Building Slide 8 (Pitch 5)...")
    build_slide_8(prs.slides[7])
    
    print("Building Slide 9 (Pitch 6)...")
    build_slide_9(prs.slides[8])
    
    print("Building Slide 10 (Pitch 7)...")
    build_slide_10(prs.slides[9])
    
    print("Building Slide 11 (Pitch 8)...")
    build_slide_11(prs.slides[10])
    
    print("Building Slide 12 (Pitch 9)...")
    build_slide_12(prs.slides[11])
    
    print("Building Slide 13 (Pitch 10)...")
    build_slide_13(prs.slides[12])
    
    # Slides 14, 15, 16: Anexe, Q&A, Checklist
    print("Building Slide 14 (Anexe)...")
    build_slide_14(prs.slides[13])
    
    print("Building Slide 15 (Q&A)...")
    build_slide_15(prs.slides[14])
    
    print("Building Slide 16 (Checklist)...")
    build_slide_16(prs.slides[15])
    
    prs.save(prs_path)
    print("Presentation saved successfully!")

if __name__ == '__main__':
    main()
