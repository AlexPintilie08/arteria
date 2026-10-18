from pathlib import Path
import json,hashlib,zipfile,argparse
from lxml import etree
from pptx import Presentation
from PIL import Image,ImageDraw
R=Path(__file__).resolve().parents[1]
a=argparse.ArgumentParser();a.add_argument('--deck',type=Path,default=R/'Template-Pitch-Hackathon.pptx');deck=a.parse_args().deck
p=Presentation(deck);read=lambda f:json.loads((R/'reports'/f).read_text(encoding='utf-8-sig'))
bounds=read('text-bounds.json');plan=read('motion-plan.json');com=read('motion-com.json')
overflow=[x for x in bounds if x['boundHeight']>x['height']+1 or x['boundWidth']>x['width']+1]
def walk(shapes):
 for sh in shapes:
  yield sh
  if sh.shape_type==6:yield from walk(sh.shapes)
manifest=[];out=[];counts={'groups':0,'textboxes':0,'tables':0,'pictures':0,'connectors':0,'notes':0}
for i,s in enumerate(p.slides,1):
 texts=[];counts['notes']+=bool(s.notes_slide.notes_text_frame.text.strip())
 for sh in walk(s.shapes):
  if sh.has_text_frame:counts['textboxes']+=1;texts.append(sh.text)
  if sh.has_table:counts['tables']+=1;texts.extend(c.text for row in sh.table.rows for c in row.cells)
  for k,n in [('groups',6),('pictures',13),('connectors',9)]:counts[k]+=sh.shape_type==n
  if sh.left<0 or sh.top<0 or sh.left+sh.width>p.slide_width+10 or sh.top+sh.height>p.slide_height+10:out.append([i,sh.name])
 manifest.append({'slide':i,'text':texts})
assert len(p.slides)==13 and counts['notes']==13
for n in ['Iustin Fagateanu','Alex Pintilie','Fagateanu Eusebiu','Backend']:assert n in '\n'.join(manifest[8]['text'])
assert len([s for s in p.slides[8].shapes if s.name.startswith('reveal:s09_member')])==3
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'};motion=[]
with zipfile.ZipFile(deck) as z:
 assert z.testzip() is None
 for i in range(1,14):
  x=etree.fromstring(z.read(f'ppt/slides/slide{i}.xml'));tr=x.xpath('.//p:transition',namespaces=ns)
  assert tr and all(t.get('advTm') in (None,'0') and t.get('advClick','1')=='1' and t.find('p:fade',ns) is not None for t in tr)
  ids=set(x.xpath('.//p:cNvPr/@id',namespaces=ns));targets=x.xpath('.//p:spTgt/@spid',namespaces=ns);assert targets and all(t in ids for t in targets)
  fx=x.xpath('.//p:cTn[@presetClass="entr"]',namespaces=ns);expected=[r for r in plan if r['slide']==i];actual=com['slides'][i-1]
  assert len(fx)==len(expected)==len(actual['effects'])
  clicks=sum(r['trigger']=='click' for r in expected);assert clicks<=4
  assert actual['transition']==3849 and abs(actual['transitionDuration']-.55)<.01 and actual['advanceOnTime']==0
  for a,b in zip(expected,actual['effects']):assert a['shape']==b['shape'] and abs(a['duration']-b['duration'])<.01 and b['effectId'] in (10,22) and b['trigger']=={'click':1,'auto':2,'after':3}[a['trigger']]
  motion.append({'slide':i,'effects':len(fx),'clicks':clicks,'targetsValid':True})
files=sorted((R/'preview').glob('Slide*.PNG'),key=lambda f:int(f.stem[5:]));assert len(files)==13
for j in range(0,13,4):
 sheet=Image.new('RGB',(1600,940),'white')
 for k,f in enumerate(files[j:j+4]):
  im=Image.open(f);assert im.size==(1920,1080);im.thumbnail((800,450));x=k%2*800;y=k//2*470;sheet.paste(im,(x,y));ImageDraw.Draw(sheet).text((x+10,y+450),f.stem,fill='black')
 sheet.save(R/'preview'/f'contact-{j//4+1}.jpg')
hash=hashlib.sha256((R/'archive/original-backup/Template-Pitch-Hackathon.backup.pptx').read_bytes()).hexdigest();assert hash=='56133b4d2fe6946ff4c5a246e4d0f8be0baa0ac2f6638a6e25f969994695d170'
assert not overflow and not out,(overflow,out)
result={'slides':13,'objects':counts,'rendered_slides':13,'text_bounds_checked':len(bounds),'text_overflow':overflow,'out_of_canvas':out,'backup_unchanged':True,'backup_sha256':hash,'pptx_sha256':hashlib.sha256(deck.read_bytes()).hexdigest(),'motion':motion,'effects':sum(x['effects'] for x in motion),'all_checks_passed':True}
for name,value in [('qa-results.json',result),('deck-manifest.json',manifest)]:(R/'reports'/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
